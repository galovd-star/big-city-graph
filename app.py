import streamlit as st
import pandas as pd
import streamlit.components.v1 as components
from pyvis.network import Network
from streamlit_gsheets import GSheetsConnection
import plotly.express as px
import plotly.graph_objects as go
import json
import base64
from io import BytesIO
from PIL import Image
import html
import re
import docx

# --- НАСТРОЙКИ СТРАНИЦЫ ---
st.set_page_config(layout="wide", page_title="БИГ СИТИ: Экосистема Вселенной", initial_sidebar_state="expanded")

# 🔗 ВСТАВЬТЕ СЮДА ССЫЛКУ НА ВАШУ GOOGLE ТАБЛИЦУ (внутри кавычек)
SHEET_URL = "https://docs.google.com/spreadsheets/d/1MyQ3vqstgSAarY9NtsAl1f1tnlbgmi-2iljHMf97zBM/edit?gid=0#gid=0"

if SHEET_URL == "ВСТАВЬТЕ_ССЫЛКУ_СЮДА":
    st.error("⚠️ Откройте код app.py на GitHub и вставьте настоящую ссылку на вашу Google Таблицу в переменную SHEET_URL.")
    st.stop()

# --- КАСТОМНЫЙ ДИЗАЙН, СТИЛИ КАРТОЧЕК И ВСПЛЫВАЮЩИХ ИМЁН В ТЕКСТЕ ---
st.markdown("""
<style>
    [data-testid="stAppViewContainer"] { background-color: #121212; }
    [data-testid="stSidebar"] { background-color: #1a1a24 !important; border-right: 1px solid #333344; }
    [data-testid="stHeader"] { background-color: rgba(0,0,0,0) !important; }
    h1, h2, h3 { color: #00d4ff !important; font-family: 'Courier New', Courier, monospace; }
    .stMarkdown p, .stText { color: #dcdcdc !important; font-size: 16px; }
    
    /* Стили для карточек галереи */
    .game-card {
        background-color: #1b1b26;
        border: 1px solid #33334d;
        border-radius: 12px;
        padding: 14px;
        margin-bottom: 8px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.5);
    }
    .game-card-header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px; }
    .game-card-type { font-size: 11px; padding: 2px 8px; border-radius: 4px; background-color: #2b2b3d; color: #00d4ff; text-transform: uppercase; font-weight: bold; }
    .game-card-title { color: #ffffff; font-size: 22px; font-weight: bold; font-family: 'Courier New', monospace; letter-spacing: 1px; margin-bottom: 6px; }
    
    /* Сетка показателей С, Л, М, С */
    .stats-grid {
        display: grid;
        grid-template-columns: 1fr 1fr;
        gap: 6px;
        background-color: #232334;
        padding: 8px 10px;
        border-radius: 8px;
        margin: 8px 0;
        font-family: 'Courier New', monospace;
        font-size: 14px;
        border-left: 3px solid #00d4ff;
    }
    .stat-item { color: #e0e0e0; }
    .stat-letter { color: #00d4ff; font-weight: bold; margin-right: 4px; }
    
    /* Плашка класса снизу карточки */
    .alignment-badge {
        text-align: center;
        padding: 5px;
        border-radius: 6px;
        font-weight: bold;
        font-size: 13px;
        text-transform: uppercase;
        letter-spacing: 1px;
        margin-top: 8px;
        margin-bottom: 6px;
    }
    .align-hero { background-color: rgba(77, 166, 255, 0.2); color: #4da6ff; border: 1px solid #4da6ff; }
    .align-villain { background-color: rgba(255, 77, 77, 0.2); color: #ff4d4d; border: 1px solid #ff4d4d; }
    .align-neutral { background-color: rgba(255, 204, 0, 0.2); color: #ffcc00; border: 1px solid #ffcc00; }

    /* Плашки связей и баланса */
    .rel-pill {
        padding: 5px 10px;
        border-radius: 6px;
        margin-bottom: 5px;
        font-size: 13px;
        display: flex;
        justify-content: space-between;
        align-items: center;
    }
    .rel-ally { background-color: rgba(77, 166, 255, 0.15); border-left: 3px solid #4da6ff; color: #d9ecff; }
    .rel-enemy { background-color: rgba(255, 77, 77, 0.15); border-left: 3px solid #ff4d4d; color: #ffd9d9; }
    .rel-neutral { background-color: rgba(160, 160, 160, 0.15); border-left: 3px solid #999999; color: #e0e0e0; }
    
    .balance-box {
        background-color: #181824;
        border: 1px solid #2e2e48;
        border-radius: 10px;
        padding: 12px;
        margin-top: 8px;
        margin-bottom: 12px;
    }
    .neighbor-card {
        background-color: #222233;
        border-radius: 6px;
        padding: 6px 10px;
        margin-bottom: 4px;
        font-size: 12px;
        border-left: 3px solid #ffcc00;
        display: flex;
        justify-content: space-between;
    }

    /* Стили для скрытых заголовков и подзаголовков в Сценариях */
    details.lvl1 { background-color: #1b1b26; border-left: 4px solid #00d4ff; border-radius: 6px; padding: 10px 15px; margin: 12px 0; overflow: visible; }
    details.lvl1 > summary { font-size: 20px; font-weight: bold; color: #00d4ff; cursor: pointer; outline: none; }
    details.lvl2 { background-color: #232332; border-left: 3px solid #ffcc00; border-radius: 6px; padding: 8px 14px; margin: 10px 0 10px 15px; overflow: visible; }
    details.lvl2 > summary { font-size: 17px; font-weight: bold; color: #ffcc00; cursor: pointer; outline: none; }
    details.lvl3 { background-color: #2b2b3c; border-left: 3px solid #ff66b2; border-radius: 6px; padding: 6px 12px; margin: 8px 0 8px 25px; overflow: visible; }
    details.lvl3 > summary { font-size: 15px; font-weight: bold; color: #ff66b2; cursor: pointer; outline: none; }
    .script-body { padding-top: 8px; color: #e0e0e0; line-height: 1.75; overflow: visible; }

    /* =========================================================
       ИНТЕРАКТИВНЫЕ УПОМИНАНИЯ ПЕРСОНАЖЕЙ ПРЯМО В ТЕКСТЕ
       ========================================================= */
    .char-mention {
        position: relative;
        display: inline-block;
        color: #00d4ff;
        background-color: rgba(0, 212, 255, 0.14);
        border-bottom: 1px dashed #00d4ff;
        border-radius: 4px;
        padding: 0px 6px;
        margin: 0 1px;
        font-weight: 600;
        cursor: pointer;
        outline: none;
        transition: background-color 0.2s ease;
    }
    .char-mention:hover, .char-mention:focus, .char-mention:active {
        background-color: rgba(0, 212, 255, 0.32);
        color: #ffffff;
        z-index: 9999;
    }
    .char-inline-popup {
        visibility: hidden;
        opacity: 0;
        position: absolute;
        left: 0;
        top: calc(100% + 6px);
        width: 265px;
        background: rgba(22, 22, 34, 0.98);
        border: 1px solid #00d4ff;
        border-radius: 12px;
        padding: 12px;
        color: #ffffff;
        font-family: 'Segoe UI', Tahoma, sans-serif;
        font-weight: normal;
        box-shadow: 0 10px 28px rgba(0,0,0,0.9), 0 0 14px rgba(0, 212, 255, 0.3);
        z-index: 99999;
        pointer-events: none;
        transition: opacity 0.18s ease, visibility 0.18s ease;
        text-align: left;
        line-height: 1.35;
    }
    .char-mention:hover .char-inline-popup,
    .char-mention:focus .char-inline-popup,
    .char-mention:active .char-inline-popup {
        visibility: visible;
        opacity: 1;
    }
    .cip-header { display: flex; justify-content: space-between; align-items: center; font-size: 11px; margin-bottom: 4px; }
    .cip-faction { color: #ffcc00; font-weight: bold; text-transform: uppercase; }
    .cip-type { background: #2b2b3d; color: #00d4ff; padding: 1px 6px; border-radius: 4px; font-weight: bold; }
    .cip-title { display: block; font-family: 'Courier New', monospace; font-size: 17px; font-weight: bold; color: #ffffff; margin-bottom: 6px; border-bottom: 1px solid #33334d; padding-bottom: 4px; }
    .cip-img-wrap { display: block; width: 100%; max-height: 175px; overflow: hidden; border-radius: 8px; margin-bottom: 7px; background: #12121a; }
    .cip-img-wrap img { width: 100%; height: 175px; object-fit: cover; display: block; }
    .cip-stats { display: grid; grid-template-columns: 1fr 1fr; gap: 4px 8px; background: #232334; padding: 6px 8px; border-radius: 6px; font-family: 'Courier New', monospace; font-size: 12px; border-left: 3px solid #00d4ff; margin-bottom: 6px; }
    .cip-stat-let { color: #00d4ff; font-weight: bold; margin-right: 4px; }
    .cip-align { display: block; text-align: center; padding: 3px; border-radius: 5px; font-weight: bold; font-size: 11px; text-transform: uppercase; letter-spacing: 1px; margin-bottom: 6px; }
    .cip-hero { background: rgba(77, 166, 255, 0.2); color: #4da6ff; border: 1px solid #4da6ff; }
    .cip-villain { background: rgba(255, 77, 77, 0.2); color: #ff4d4d; border: 1px solid #ff4d4d; }
    .cip-neutral { background: rgba(255, 204, 0, 0.2); color: #ffcc00; border: 1px solid #ffcc00; }
    .cip-desc { display: block; font-size: 12px; color: #cccccc; margin-bottom: 5px; }
    .cip-bonuses { display: block; font-size: 11px; color: #ff99cc; margin-bottom: 5px; }
    .cip-lore { display: block; font-size: 11px; color: #aaaaaa; border-top: 1px solid #2e2e42; padding-top: 4px; margin-top: 4px; }
</style>
""", unsafe_allow_html=True)

PALETTE = ["#ff4d4d", "#4da6ff", "#9933ff", "#33cc33", "#ffcc00", "#ff66b2", "#339999", "#ff9933", "#996633", "#6699ff"]
ALIGN_COLORS = {"Герой": "#4da6ff", "Злодей": "#ff4d4d", "Нейтралитет": "#ffcc00"}

# --- ПОДКЛЮЧЕНИЕ К GOOGLE SHEETS ---
conn = st.connection("gsheets", type=GSheetsConnection)

CARD_COLS = [
    'Name', 'Card_Type', 'Alignment', 'Faction',
    'Strength', 'Agility', 'Mastery', 'Wit',
    'Description', 'Image_Data', 'Linked_Character'
]
LOC_COLS = [
    'Name', 'Region', 'Loc_Type', 'Controlling_Faction',
    'X', 'Y', 'Description', 'Map_Image', 'Linked_Characters', 'Linked_Scripts'
]
TIME_COLS = [
    'Order_Index', 'Date_Label', 'Era', 'Title',
    'Description', 'Participants', 'Location', 'Linked_Script'
]
FAC_COLS = [
    'Name', 'Org_Category', 'Motto', 'Leader', 'Ideology', 'Base_Location',
    'Influence', 'Color', 'Description', 'Logo_Data', 'Allies', 'Enemies'
]

NUMERIC_COLS = {'Strength', 'Agility', 'Mastery', 'Wit', 'Order_Index', 'X', 'Y', 'Influence', 'Weight'}

def clean_text_cell(val):
    if pd.isna(val):
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val)

def safe_read_sheet(worksheet_name, default_cols):
    try:
        df = conn.read(spreadsheet=SHEET_URL, worksheet=worksheet_name)
        df = df.dropna(how='all')
        if df.empty or len(df.columns) == 0:
            return pd.DataFrame(columns=default_cols)
        for col in default_cols:
            if col not in df.columns:
                if col in ['Strength', 'Agility', 'Mastery', 'Wit', 'Order_Index']:
                    df[col] = 0
                elif col in ['X', 'Y', 'Influence']:
                    df[col] = 50
                elif col == 'Card_Type':
                    df[col] = 'Персонаж'
                elif col == 'Alignment':
                    df[col] = 'Нейтралитет'
                elif col == 'Org_Category':
                    df[col] = 'Группировка'
                elif col == 'Color':
                    df[col] = '#00d4ff'
                elif col in [
                    "Content", "Image_Data", "Заметки", "Description", "Linked_Character",
                    "Map_Image", "Linked_Characters", "Linked_Scripts", "Participants",
                    "Location", "Linked_Script", "Motto", "Leader", "Ideology",
                    "Base_Location", "Logo_Data", "Allies", "Enemies"
                ]:
                    df[col] = ""
                else:
                    df[col] = "Не указано"
        
        df = df.fillna("")
        # Принудительно переводим все текстовые столбцы (включая имена с цифрами) в строки
        for col in df.columns:
            if col not in NUMERIC_COLS:
                df[col] = df[col].map(clean_text_cell)
        return df
    except Exception as e:
        st.error(f"⚠️ Ошибка загрузки листа '{worksheet_name}'. База защищена от стирания. Текст ошибки: {e}")
        st.stop()

@st.cache_data(ttl=3600)
def load_all_data():
    nodes = safe_read_sheet("Nodes", ['Имя / Название персонажа', 'Категория плана', 'Пол', 'Группировка', 'Заметки']).replace("", "Не указано")
    edges = safe_read_sheet("Edges", ['Source', 'Target', 'Weight', 'Type'])
    archives = safe_read_sheet("Archives", ['Title', 'Category', 'Content'])
    cards = safe_read_sheet("Cards", CARD_COLS)
    for stat_col in ['Strength', 'Agility', 'Mastery', 'Wit']:
        cards[stat_col] = pd.to_numeric(cards[stat_col], errors='coerce').fillna(0).astype(int)
    scripts = safe_read_sheet("Scripts", ['Title', 'Act', 'Status', 'Content'])
    
    locations = safe_read_sheet("Locations", LOC_COLS)
    for c_xy in ['X', 'Y']:
        locations[c_xy] = pd.to_numeric(locations[c_xy], errors='coerce').fillna(50.0)
        
    timeline = safe_read_sheet("Timeline", TIME_COLS)
    timeline['Order_Index'] = pd.to_numeric(timeline['Order_Index'], errors='coerce').fillna(1).astype(int)
    
    factions = safe_read_sheet("Factions", FAC_COLS)
    factions['Influence'] = pd.to_numeric(factions['Influence'], errors='coerce').fillna(50).astype(int)
    if 'Org_Category' not in factions.columns:
        factions['Org_Category'] = 'Группировка'
    factions['Org_Category'] = factions['Org_Category'].replace("", "Группировка")
    
    return nodes, edges, archives, cards, scripts, locations, timeline, factions

@st.cache_data(ttl=3600)
def load_config():
    try:
        cfg = conn.read(spreadsheet=SHEET_URL, worksheet="Config")
        if not cfg.empty and 'Settings' in cfg.columns:
            val = cfg['Settings'].iloc[0]
            if pd.notna(val) and str(val).strip() != "":
                return json.loads(val)
    except Exception:
        pass
    return {"colors": {}, "options": {}, "region_maps": {}}

def save_sheet(worksheet_name, df):
    try:
        conn.update(spreadsheet=SHEET_URL, worksheet=worksheet_name, data=df)
        st.cache_data.clear()
    except Exception as e:
        st.error(f"⚠️ Ошибка сохранения в лист '{worksheet_name}'. Убедитесь, что лист '{worksheet_name}' создан в вашей Google Таблице! Детали: {e}")

def save_config_to_cloud(config_dict):
    df = pd.DataFrame({"Settings": [json.dumps(config_dict, ensure_ascii=False)]})
    save_sheet("Config", df)

# --- ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ СВЯЗЕЙ И КАРТОЧЕК ---
def get_character_relations(char_name, edges_df):
    if not char_name or edges_df.empty:
        return []
    char_clean = str(char_name).strip()
    rels = []
    for _, row in edges_df.iterrows():
        src = str(row.get('Source', '')).strip()
        tgt = str(row.get('Target', '')).strip()
        rtype = str(row.get('Type', 'Связь')).strip()
        try:
            w = int(float(row.get('Weight', 1)))
        except ValueError:
            w = 1
        if src == char_clean:
            rels.append({"partner": tgt, "type": rtype, "weight": w, "dir": "→"})
        elif tgt == char_clean:
            rels.append({"partner": src, "type": rtype, "weight": w, "dir": "←"})
    return rels

def resolve_linked_character(card_row, valid_chars_list):
    explicit_link = str(card_row.get('Linked_Character', '')).strip()
    if explicit_link and explicit_link != "— Нет привязки —" and explicit_link in valid_chars_list:
        return explicit_link
    card_name = str(card_row.get('Name', '')).strip()
    for ch in valid_chars_list:
        # Исправленная строка: принудительно делаем ch строкой перед .strip()
        if str(ch).strip().lower() == card_name.lower():
            return ch
    return None

def get_cards_for_character(char_name, cards_df):
    if cards_df.empty or not char_name:
        return pd.DataFrame(columns=CARD_COLS)
    clean_name = str(char_name).strip().lower()
    mask = (
        (cards_df['Linked_Character'].astype(str).str.strip().str.lower() == clean_name) |
        (cards_df['Name'].astype(str).str.strip().str.lower() == clean_name)
    )
    return cards_df[mask]

# --- СБОРКА ПОЛНОГО СЛОВАРЯ ДАННЫХ ВСЕХ ПЕРСОНАЖЕЙ (ДЛЯ ГРАФА И ТЕКСТОВ) ---
def build_universe_character_lookup(nodes_df, edges_df, cards_df):
    lookup = {}
    color_options = [c for c in nodes_df.columns if c not in ["Имя / Название персонажа", "Заметки"]]

    for _, row in nodes_df.iterrows():
        name = str(row.get('Имя / Название персонажа', '')).strip()
        if not name or name == "Не указано":
            continue
        lore_lines = [f"{c}: {row.get(c, 'Не указано')}" for c in color_options if str(row.get(c, 'Не указано')) != 'Не указано']
        if str(row.get('Заметки', '')).strip() not in ['', 'Не указано']:
            lore_lines.append(f"Лор: {row.get('Заметки')}")

        rels = get_character_relations(name, edges_df)
        rel_strings = [f"{r['type']} с {r['partner']} ({'⭐' * r['weight']})" for r in rels]

        char_cards = get_cards_for_character(name, cards_df)
        main_card = None
        bonus_names = []
        if not char_cards.empty:
            pers_cards = char_cards[char_cards['Card_Type'] == 'Персонаж']
            main_card = pers_cards.iloc[0] if not pers_cards.empty else char_cards.iloc[0]
            bonus_cards = char_cards[char_cards['Card_Type'] == 'Бонус']
            bonus_names = bonus_cards['Name'].tolist()

        if main_card is not None:
            lookup[name] = {
                "has_card": True,
                "card_name": str(main_card.get('Name', name)),
                "card_type": str(main_card.get('Card_Type', 'Персонаж')),
                "alignment": str(main_card.get('Alignment', 'Нейтралитет')),
                "faction": str(main_card.get('Faction', '')),
                "str": int(main_card.get('Strength', 0)),
                "agi": int(main_card.get('Agility', 0)),
                "mas": int(main_card.get('Mastery', 0)),
                "wit": int(main_card.get('Wit', 0)),
                "desc": str(main_card.get('Description', '')),
                "image": str(main_card.get('Image_Data', '')).strip(),
                "bonuses": bonus_names,
                "lore": lore_lines,
                "relations": rel_strings,
                "raw_rels": rels
            }
        else:
            lookup[name] = {
                "has_card": False,
                "card_name": name,
                "lore": lore_lines,
                "relations": rel_strings,
                "raw_rels": rels
            }

    for _, crd in cards_df.iterrows():
        c_name = str(crd.get('Name', '')).strip()
        if c_name and c_name not in lookup and str(crd.get('Card_Type', '')) == 'Персонаж':
            lookup[c_name] = {
                "has_card": True,
                "card_name": c_name,
                "card_type": "Персонаж",
                "alignment": str(crd.get('Alignment', 'Нейтралитет')),
                "faction": str(crd.get('Faction', '')),
                "str": int(crd.get('Strength', 0)),
                "agi": int(crd.get('Agility', 0)),
                "mas": int(crd.get('Mastery', 0)),
                "wit": int(crd.get('Wit', 0)),
                "desc": str(crd.get('Description', '')),
                "image": str(crd.get('Image_Data', '')).strip(),
                "bonuses": [],
                "lore": [],
                "relations": [],
                "raw_rels": []
            }
    return lookup

# --- УМНЫЙ ПОИСК И ПОДСВЕТКА ИМЁН ПЕРСОНАЖЕЙ В ТЕКСТЕ (С УЧЁТОМ ПАДЕЖЕЙ) ---
def get_russian_stem_pattern(w, is_multiword=False):
    if not re.match(r'^[А-Яа-яЁё]+$', w):
        return re.escape(w)
    w_low = w.lower()
    if len(w) >= 5 and w_low.endswith(('ая', 'яя', 'ый', 'ий', 'ой', 'ое', 'ее')):
        return re.escape(w[:-2]) + r'[А-Яа-яЁё]{0,4}'
    if len(w) >= 4 or (is_multiword and len(w) >= 3):
        stem = w[:-1] if w_low[-1] in 'аяеёиоуыэюьй' else w
        return re.escape(stem) + r'[А-Яа-яЁё]{0,3}'
    return re.escape(w)

def build_character_matchers(char_names):
    alias_to_canon = {}
    for name in char_names:
        clean = str(name).strip()
        if not clean or clean == "Не указано":
            continue
        if clean.isdigit():
            continue
        alias_to_canon[clean] = clean
        m = re.match(r'^(.+?)\s*\((.+?)\)$', clean)
        if m:
            part1, part2 = m.group(1).strip(), m.group(2).strip()
            if len(part1) >= 3:
                alias_to_canon[part1] = clean
            if len(part2) >= 3:
                alias_to_canon[part2] = clean

    sorted_aliases = sorted(alias_to_canon.keys(), key=len, reverse=True)
    patterns = []
    for alias in sorted_aliases:
        words = alias.split()
        is_multi = len(words) > 1
        word_pats = [get_russian_stem_pattern(w, is_multi) for w in words]
        full_pat = r'(?<![А-Яа-яЁёA-Za-z0-9_])(' + r'\s+'.join(word_pats) + r')(?![А-Яа-яЁёA-Za-z0-9_])'
        patterns.append((re.compile(full_pat, re.IGNORECASE), alias_to_canon[alias]))
    return patterns

def find_character_mentions_across_ecosystem(char_name, matchers, archives_df, scripts_df, locations_df, timeline_df):
    char_pats = [pat for pat, canon in matchers if canon == char_name]
    results = {"archives": [], "scripts": [], "locations": [], "timeline": []}
    if not char_pats:
        return results

    def has_match(text):
        t = str(text)
        return any(p.search(t) for p in char_pats)

    def extract_snippet(text):
        t = str(text)
        for p in char_pats:
            m = p.search(t)
            if m:
                start = max(0, m.start() - 55)
                end = min(len(t), m.end() + 55)
                return ("..." if start > 0 else "") + t[start:end].replace("\n", " ") + ("..." if end < len(t) else "")
        return ""

    for _, r in archives_df.iterrows():
        if has_match(r.get('Title', '')) or has_match(r.get('Content', '')):
            results["archives"].append({
                "title": r.get('Title', 'Без названия'),
                "cat": r.get('Category', 'Общее'),
                "snippet": extract_snippet(r.get('Content', ''))
            })

    for _, r in scripts_df.iterrows():
        if has_match(r.get('Title', '')) or has_match(r.get('Content', '')):
            results["scripts"].append({
                "title": r.get('Title', 'Без названия'),
                "act": r.get('Act', ''),
                "status": r.get('Status', ''),
                "snippet": extract_snippet(r.get('Content', ''))
            })

    for _, r in locations_df.iterrows():
        linked_list = [x.strip() for x in str(r.get('Linked_Characters', '')).split(',') if x.strip()]
        if char_name in linked_list or has_match(r.get('Description', '')) or has_match(r.get('Name', '')):
            results["locations"].append({
                "name": r.get('Name', ''),
                "region": r.get('Region', ''),
                "type": r.get('Loc_Type', '')
            })

    for _, r in timeline_df.iterrows():
        part_list = [x.strip() for x in str(r.get('Participants', '')).split(',') if x.strip()]
        if char_name in part_list or has_match(r.get('Description', '')) or has_match(r.get('Title', '')):
            results["timeline"].append({
                "date": r.get('Date_Label', ''),
                "era": r.get('Era', ''),
                "title": r.get('Title', '')
            })

    return results

def build_inline_popup_span(char_name, info):
    parts = []
    if info.get("has_card"):
        align = info.get("alignment", "Нейтралитет")
        align_cls = "cip-hero" if align == "Герой" else ("cip-villain" if align == "Злодей" else "cip-neutral")
        parts.append(
            f'<span class="cip-header">'
            f'<span class="cip-faction">{html.escape(str(info.get("faction", "")))}</span>'
            f'<span class="cip-type">{html.escape(str(info.get("card_type", "Персонаж")))}</span>'
            f'</span>'
        )
        parts.append(f'<span class="cip-title">{html.escape(str(info.get("card_name", char_name)))}</span>')
        img_src = info.get("image", "")
        if img_src:
            parts.append(f'<span class="cip-img-wrap"><img src="{img_src}" alt="card"/></span>')
        parts.append(
            f'<span class="cip-stats">'
            f'<span><span class="cip-stat-let">С</span>{info.get("str", 0)}</span>'
            f'<span><span class="cip-stat-let">Л</span>{info.get("agi", 0)}</span>'
            f'<span><span class="cip-stat-let">М</span>{info.get("mas", 0)}</span>'
            f'<span><span class="cip-stat-let">С</span>{info.get("wit", 0)}</span>'
            f'</span>'
        )
        parts.append(f'<span class="cip-align {align_cls}">{html.escape(str(align))}</span>')
        if info.get("desc"):
            parts.append(f'<span class="cip-desc">{html.escape(str(info.get("desc")))}</span>')
        if info.get("bonuses"):
            parts.append(f'<span class="cip-bonuses">🎁 Бонусы: {html.escape(", ".join(info["bonuses"]))}</span>')
    else:
        parts.append(f'<span class="cip-title">👤 {html.escape(char_name)}</span>')

    if info.get("lore"):
        lore_html = "<br>".join(f"▪ {html.escape(str(x))}" for x in info["lore"])
        parts.append(f'<span class="cip-lore">{lore_html}</span>')

    if info.get("relations"):
        rel_html = "<br>".join(f"🔗 {html.escape(str(r))}" for r in info["relations"][:4])
        parts.append(f'<span class="cip-lore">{rel_html}</span>')

    return "".join(parts)

def highlight_characters_in_escaped_line(escaped_line, matchers, popup_span_map, found_chars_list):
    replacements = []

    def make_replacer(canon_name):
        def _repl(match_obj):
            surface = match_obj.group(0)
            idx = len(replacements)
            replacements.append((surface, canon_name))
            if canon_name not in found_chars_list:
                found_chars_list.append(canon_name)
            return f"__BIGCITY_CHAR_{idx}__"
        return _repl

    work_line = escaped_line
    for pat, canon in matchers:
        if canon in popup_span_map:
            work_line = pat.sub(make_replacer(canon), work_line)

    for idx, (surface, canon) in enumerate(replacements):
        popup_inner = popup_span_map[canon]
        badge_html = (
            f'<span class="char-mention" tabindex="0">'
            f'👤 {surface}'
            f'<span class="char-inline-popup">{popup_inner}</span>'
            f'</span>'
        )
        work_line = work_line.replace(f"__BIGCITY_CHAR_{idx}__", badge_html)

    return work_line

def render_character_popover_bar(found_chars, char_lookup, key_prefix="Mention"):
    if not found_chars:
        return
    st.caption("👥 **Упомянутые в тексте персонажи** *(нажмите на кнопку для открытия полного досье и карточки)*:")
    cols = st.columns(min(len(found_chars), 4))
    for i, ch_name in enumerate(found_chars):
        info = char_lookup.get(ch_name, {})
        with cols[i % 4]:
            with st.popover(f"👤 {ch_name}", use_container_width=True):
                if info.get("has_card"):
                    al_str = info.get("alignment", "Нейтралитет")
                    al_cls = "align-hero" if al_str == "Герой" else ("align-villain" if al_str == "Злодей" else "align-neutral")
                    st.markdown(f"""
                    <div class="game-card">
                        <div class="game-card-header">
                            <span style="color:#ffcc00; font-size:12px;">{html.escape(str(info.get('faction', '')))}</span>
                            <span class="game-card-type">{html.escape(str(info.get('card_type', 'Персонаж')))}</span>
                        </div>
                        <div class="game-card-title">{html.escape(str(info.get('card_name', ch_name)))}</div>
                    </div>
                    """, unsafe_allow_html=True)
                    if info.get("image"):
                        st.image(info["image"], use_container_width=True)
                    st.markdown(f"""
                    <div class="stats-grid">
                        <div class="stat-item"><span class="stat-letter">С</span>{info.get('str', 0)}</div>
                        <div class="stat-item"><span class="stat-letter">Л</span>{info.get('agi', 0)}</div>
                        <div class="stat-item"><span class="stat-letter">М</span>{info.get('mas', 0)}</div>
                        <div class="stat-item"><span class="stat-letter">С</span>{info.get('wit', 0)}</div>
                    </div>
                    <div class="alignment-badge {al_cls}">{html.escape(al_str)}</div>
                    """, unsafe_allow_html=True)
                    if info.get("desc"):
                        st.caption(info["desc"])
                    if info.get("bonuses"):
                        st.markdown(f"🎁 **Бонусы:** {', '.join(info['bonuses'])}")
                else:
                    st.markdown(f"### 👤 {ch_name}")

                if info.get("lore"):
                    st.write("---")
                    for l_item in info["lore"]:
                        st.markdown(f"▪ {l_item}")

                if info.get("raw_rels"):
                    st.write("---")
                    st.markdown("**🔗 Связи на графе:**")
                    for r in info["raw_rels"]:
                        t_low = r['type'].lower()
                        css_c = "rel-enemy" if t_low in ['вражда', 'конфликт'] else ("rel-ally" if t_low in ['союз', 'родство', 'подчинение'] else "rel-neutral")
                        st.markdown(f"""
                        <div class="rel-pill {css_c}">
                            <span><b>{html.escape(r['partner'])}</b> ({html.escape(r['type'])})</span>
                            <span>{'⭐' * r['weight']}</span>
                        </div>
                        """, unsafe_allow_html=True)

# --- ЧТЕНИЕ И СЖАТИЕ ФАЙЛОВ ---
def read_uploaded_document(uploaded_file):
    filename = uploaded_file.name
    base_title = filename.rsplit('.', 1)[0]
    if filename.lower().endswith('.docx'):
        doc = docx.Document(uploaded_file)
        lines = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                lines.append("")
                continue
            style_name = p.style.name.lower() if p.style else ""
            if 'heading 1' in style_name or 'заголовок 1' in style_name:
                lines.append(f"# {text}")
            elif 'heading 2' in style_name or 'заголовок 2' in style_name:
                lines.append(f"## {text}")
            elif 'heading 3' in style_name or 'заголовок 3' in style_name:
                lines.append(f"### {text}")
            else:
                lines.append(p.text)
        return base_title, "\n".join(lines)
    else:
        raw_bytes = uploaded_file.read()
        try:
            content = raw_bytes.decode('utf-8')
        except UnicodeDecodeError:
            content = raw_bytes.decode('cp1251', errors='replace')
        return base_title, content

def compress_uploaded_image(uploaded_file, max_size=(420, 580), quality=72):
    img = Image.open(uploaded_file)
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    img.thumbnail(max_size)
    buffer = BytesIO()
    img.save(buffer, format="JPEG", quality=quality, optimize=True)
    b64_str = base64.b64encode(buffer.getvalue()).decode("utf-8")
    return f"data:image/jpeg;base64,{b64_str}"

# --- ФУНКЦИЯ ПОСТРОЕНИЯ ЧИСЛОВОЙ ПРЯМОЙ (ОСИ ПОКАЗАТЕЛЕЙ) ---
def build_stat_number_line_figure(
    df_cards, stat_col, stat_label,
    highlight_val=None, highlight_name="Текущий ввод",
    tolerance=0, height=190
):
    fig = go.Figure()

    fig.add_hline(y=0, line_width=2, line_color="#444466", opacity=0.8)

    if highlight_val is not None and tolerance > 0:
        fig.add_vrect(
            x0=highlight_val - tolerance,
            x1=highlight_val + tolerance,
            fillcolor="rgba(0, 212, 255, 0.14)",
            line_width=1,
            line_dash="dot",
            line_color="#00d4ff",
            annotation_text=f"Зона ±{tolerance} ({highlight_val - tolerance} … {highlight_val + tolerance})",
            annotation_position="top left",
            annotation_font=dict(color="#00d4ff", size=11)
        )

    if not df_cards.empty:
        df_plot = df_cards.copy()
        df_plot[stat_col] = pd.to_numeric(df_plot[stat_col], errors='coerce').fillna(0).astype(int)
        df_plot = df_plot.sort_values(by=stat_col)
        offsets = []
        counts_per_val = {}
        for val in df_plot[stat_col]:
            c = counts_per_val.get(val, 0)
            counts_per_val[val] = c + 1
            level = (c // 2 + 1) * 0.16 * (1 if c % 2 == 0 else -1) if c > 0 else 0.0
            offsets.append(level)
        df_plot['y_offset'] = offsets

        for align_val, group in df_plot.groupby('Alignment'):
            color_hex = ALIGN_COLORS.get(align_val, "#ffcc00")
            fig.add_trace(go.Scatter(
                x=group[stat_col],
                y=group['y_offset'],
                mode='markers+text',
                name=str(align_val),
                text=group['Name'],
                textposition='top center',
                textfont=dict(size=10, color='#d0d0e0'),
                marker=dict(size=12, color=color_hex, line=dict(width=1.5, color='#121212'), symbol='circle'),
                customdata=group[['Faction', 'Strength', 'Agility', 'Mastery', 'Wit', 'Total_Power']],
                hovertemplate=(
                    "<b>%{text}</b> (" + str(align_val) + " • %{customdata[0]})<br>" +
                    f"<b>{stat_label}: %{{x}}</b><br>" +
                    "С: %{customdata[1]} | Л: %{customdata[2]} | М: %{customdata[3]} | С: %{customdata[4]}<br>" +
                    "Сумма: %{customdata[5]}<extra></extra>"
                )
            ))

    if highlight_val is not None:
        fig.add_trace(go.Scatter(
            x=[highlight_val],
            y=[0],
            mode='markers+text',
            name='Редактируемая карточка',
            text=[f"★ {highlight_name} ({highlight_val})"],
            textposition='bottom center',
            textfont=dict(size=12, color='#ffd700', family='Courier New'),
            marker=dict(size=18, color='#ffd700', symbol='diamond', line=dict(width=2, color='#ffffff')),
            hovertemplate=f"<b>{highlight_name}</b><br>{stat_label}: <b>{highlight_val}</b><extra></extra>"
        ))

    fig.update_layout(
        title=dict(text=f"📏 {stat_label}", font=dict(size=14, color="#00d4ff"), x=0.01, y=0.92),
        height=height,
        margin=dict(l=20, r=20, t=35, b=25),
        paper_bgcolor='#161622',
        plot_bgcolor='#161622',
        font=dict(color='white'),
        xaxis=dict(
            title="",
            gridcolor='#26263a',
            zeroline=False,
            tickfont=dict(size=11, color='#aaaaaa')
        ),
        yaxis=dict(
            visible=False,
            range=[-0.65, 0.65]
        ),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1, font=dict(size=11))
    )
    return fig

# --- ПАРСИНГ СКРЫТЫХ ЗАГОЛОВКОВ И АВТО-ПОДСВЕТКА ИМЁН ПЕРСОНАЖЕЙ ---
def render_text_with_spoilers_and_characters(raw_text, matchers, popup_span_map, enable_spoilers=True):
    lines = str(raw_text).split('\n')
    html_out = []
    open_levels = []
    found_chars = []

    def close_until(target_level):
        while open_levels and open_levels[-1] >= target_level:
            html_out.append("</div></details>")
            open_levels.pop()

    for line in lines:
        stripped = line.strip()
        if enable_spoilers and stripped.startswith('### '):
            close_until(3)
            title_escaped = html.escape(stripped[4:].strip())
            title_hl = highlight_characters_in_escaped_line(title_escaped, matchers, popup_span_map, found_chars)
            html_out.append(f'<details class="lvl3"><summary>🔹 {title_hl}</summary><div class="script-body">')
            open_levels.append(3)
        elif enable_spoilers and stripped.startswith('## '):
            close_until(2)
            title_escaped = html.escape(stripped[3:].strip())
            title_hl = highlight_characters_in_escaped_line(title_escaped, matchers, popup_span_map, found_chars)
            html_out.append(f'<details class="lvl2"><summary>📂 {title_hl}</summary><div class="script-body">')
            open_levels.append(2)
        elif enable_spoilers and stripped.startswith('# '):
            close_until(1)
            title_escaped = html.escape(stripped[2:].strip())
            title_hl = highlight_characters_in_escaped_line(title_escaped, matchers, popup_span_map, found_chars)
            html_out.append(f'<details class="lvl1"><summary>📁 {title_hl}</summary><div class="script-body">')
            open_levels.append(1)
        else:
            escaped_line = html.escape(line)
            hl_line = highlight_characters_in_escaped_line(escaped_line, matchers, popup_span_map, found_chars)
            html_out.append(hl_line + "<br>")

    close_until(1)
    return "".join(html_out), found_chars

# --- ВСТРАИВАНИЕ ИНТЕРАКТИВНОЙ КАРТОЧКИ ПРИ НАВЕДЕНИИ В HTML ГРАФА ---
def inject_card_hover_into_graph(raw_html, node_popup_dict):
    popup_json = json.dumps(node_popup_dict, ensure_ascii=False)
    custom_injection = f"""
    <style>
        #graph-card-popup {{
            position: fixed;
            display: none;
            z-index: 99999;
            width: 265px;
            background: rgba(24, 24, 36, 0.96);
            border: 1px solid #00d4ff;
            border-radius: 12px;
            padding: 12px;
            color: #ffffff;
            font-family: 'Segoe UI', Tahoma, sans-serif;
            box-shadow: 0 8px 25px rgba(0, 0, 0, 0.85), 0 0 12px rgba(0, 212, 255, 0.25);
            pointer-events: none;
            backdrop-filter: blur(6px);
        }}
        .gcp-header {{ display: flex; justify-content: space-between; align-items: center; font-size: 11px; margin-bottom: 4px; }}
        .gcp-faction {{ color: #ffcc00; font-weight: bold; text-transform: uppercase; }}
        .gcp-type {{ background: #2b2b3d; color: #00d4ff; padding: 2px 6px; border-radius: 4px; font-weight: bold; }}
        .gcp-title {{ font-family: 'Courier New', monospace; font-size: 18px; font-weight: bold; color: #ffffff; margin-bottom: 6px; border-bottom: 1px solid #33334d; padding-bottom: 4px; }}
        .gcp-img-wrap {{ width: 100%; max-height: 195px; overflow: hidden; border-radius: 8px; margin-bottom: 8px; text-align: center; background: #12121a; }}
        .gcp-img-wrap img {{ width: 100%; height: 195px; object-fit: cover; display: block; }}
        .gcp-stats {{ display: grid; grid-template-columns: 1fr 1fr; gap: 4px 8px; background: #232334; padding: 6px 9px; border-radius: 6px; font-family: 'Courier New', monospace; font-size: 13px; border-left: 3px solid #00d4ff; margin-bottom: 6px; }}
        .gcp-stat-let {{ color: #00d4ff; font-weight: bold; margin-right: 4px; }}
        .gcp-align {{ text-align: center; padding: 4px; border-radius: 5px; font-weight: bold; font-size: 11px; text-transform: uppercase; letter-spacing: 1px; margin-bottom: 6px; }}
        .gcp-hero {{ background: rgba(77, 166, 255, 0.2); color: #4da6ff; border: 1px solid #4da6ff; }}
        .gcp-villain {{ background: rgba(255, 77, 77, 0.2); color: #ff4d4d; border: 1px solid #ff4d4d; }}
        .gcp-neutral {{ background: rgba(255, 204, 0, 0.2); color: #ffcc00; border: 1px solid #ffcc00; }}
        .gcp-desc {{ font-size: 12px; color: #cccccc; margin-bottom: 6px; line-height: 1.35; }}
        .gcp-bonuses {{ font-size: 11px; color: #ff99cc; margin-bottom: 6px; }}
        .gcp-lore {{ font-size: 11px; color: #aaaaaa; border-top: 1px solid #2e2e42; padding-top: 5px; line-height: 1.4; }}
    </style>

    <div id="graph-card-popup"></div>

    <script type="text/javascript">
        (function() {{
            const popupData = {popup_json};
            const popupEl = document.getElementById('graph-card-popup');
            let isPinned = false;

            function escapeHtml(str) {{
                if (!str) return "";
                return String(str).replace(/&/g, "&amp;").replace(/</g, "&lt;").replace(/>/g, "&gt;").replace(/"/g, "&quot;");
            }}

            function buildCardHtml(nodeId) {{
                const info = popupData[nodeId];
                if (!info) return "";

                let html = "";
                if (info.has_card) {{
                    let alignClass = "gcp-neutral";
                    if (info.alignment === "Герой") alignClass = "gcp-hero";
                    else if (info.alignment === "Злодей") alignClass = "gcp-villain";

                    html += `<div class="gcp-header">
                                <span class="gcp-faction">${{escapeHtml(info.faction)}}</span>
                                <span class="gcp-type">${{escapeHtml(info.card_type)}}</span>
                             </div>`;
                    html += `<div class="gcp-title">${{escapeHtml(info.card_name)}}</div>`;
                    if (info.image) {{
                        html += `<div class="gcp-img-wrap"><img src="${{info.image}}" alt="card"/></div>`;
                    }}
                    html += `<div class="gcp-stats">
                                <div><span class="gcp-stat-let">С</span>${{info.str}}</div>
                                <div><span class="gcp-stat-let">Л</span>${{info.agi}}</div>
                                <div><span class="gcp-stat-let">М</span>${{info.mas}}</div>
                                <div><span class="gcp-stat-let">С</span>${{info.wit}}</div>
                             </div>`;
                    html += `<div class="gcp-align ${{alignClass}}">${{escapeHtml(info.alignment)}}</div>`;
                    if (info.desc) {{
                        html += `<div class="gcp-desc">${{escapeHtml(info.desc)}}</div>`;
                    }}
                    if (info.bonuses && info.bonuses.length > 0) {{
                        html += `<div class="gcp-bonuses">🎁 <b>Бонусы:</b> ${{escapeHtml(info.bonuses.join(', '))}}</div>`;
                    }}
                }} else {{
                    html += `<div class="gcp-title">👤 ${{escapeHtml(nodeId)}}</div>`;
                }}

                if (info.lore && info.lore.length > 0) {{
                    html += `<div class="gcp-lore">` + info.lore.map(item => `▪ ${{escapeHtml(item)}}`).join('<br>') + `</div>`;
                }}
                return html;
            }}

            function positionPopup(nodeId) {{
                if (typeof network === 'undefined') return;
                const pos = network.getPositions([nodeId])[nodeId];
                if (!pos) return;
                const domPos = network.canvasToDOM(pos);
                const pad = 18;
                let left = domPos.x + 24;
                let top = domPos.y - 40;

                const rect = popupEl.getBoundingClientRect();
                const w = rect.width || 275;
                const h = rect.height || 360;

                if (left + w + pad > window.innerWidth) {{
                    left = Math.max(pad, domPos.x - w - 24);
                }}
                if (top + h + pad > window.innerHeight) {{
                    top = Math.max(pad, window.innerHeight - h - pad);
                }}
                if (top < pad) top = pad;

                popupEl.style.left = left + 'px';
                popupEl.style.top = top + 'px';
            }}

            if (typeof network !== 'undefined') {{
                network.setOptions({{ interaction: {{ hover: true, tooltipDelay: 50 }} }});

                network.on("hoverNode", function(params) {{
                    if (isPinned) return;
                    const content = buildCardHtml(params.node);
                    if (content) {{
                        popupEl.innerHTML = content;
                        popupEl.style.display = 'block';
                        positionPopup(params.node);
                    }}
                }});

                network.on("blurNode", function() {{
                    if (!isPinned) {{
                        popupEl.style.display = 'none';
                    }}
                }});

                network.on("selectNode", function(params) {{
                    if (params.nodes.length > 0) {{
                        const nodeId = params.nodes[0];
                        const content = buildCardHtml(nodeId);
                        if (content) {{
                            isPinned = true;
                            popupEl.innerHTML = content;
                            popupEl.style.display = 'block';
                            positionPopup(nodeId);
                        }}
                    }}
                }});

                network.on("deselectNode", function() {{
                    isPinned = false;
                    popupEl.style.display = 'none';
                }});

                network.on("dragStart", function() {{
                    if (!isPinned) popupEl.style.display = 'none';
                }});
            }}
        }})();
    </script>
    """
    return raw_html.replace("</body>", custom_injection + "\n</body>")

# --- ГЕНЕРАЦИЯ ЭКСПОРТА «БИБЛИИ ВСЕЛЕННОЙ» (.MD И .DOCX) ---
def generate_universe_markdown(sections, nodes_df, edges_df, cards_df, factions_df, locations_df, timeline_df, archives_df, scripts_df):
    md = ["# 🌆 БИБЛИЯ ВСЕЛЕННОЙ: БИГ СИТИ\n"]
    if "Персонажи и Связи" in sections:
        md.append("## 🌐 1. Персонажи и Карта Связей\n")
        for _, r in nodes_df.iterrows():
            name = r.get('Имя / Название персонажа', '')
            md.append(f"### 👤 {name}")
            for col in nodes_df.columns:
                if col not in ['Имя / Название персонажа', 'Заметки'] and str(r.get(col, '')) != 'Не указано':
                    md.append(f"- **{col}:** {r.get(col)}")
            if str(r.get('Заметки', '')).strip() not in ['', 'Не указано']:
                md.append(f"- **Лор / Заметки:** {r.get('Заметки')}")
            rels = get_character_relations(name, edges_df)
            if rels:
                md.append("- **Связи:** " + "; ".join([f"{x['type']} с {x['partner']} (Сила: {x['weight']})" for x in rels]))
            md.append("")

    if "Игровые Карточки" in sections and not cards_df.empty:
        md.append("## 🃏 2. Игровые Карточки и Характеристики\n")
        for _, c in cards_df.iterrows():
            tot = int(c.get('Strength', 0)) + int(c.get('Agility', 0)) + int(c.get('Mastery', 0)) + int(c.get('Wit', 0))
            md.append(f"### 🃏 {c.get('Name')} ({c.get('Card_Type')} • {c.get('Alignment')} • {c.get('Faction')})")
            md.append(f"- **Показатели:** Сила (С): {c.get('Strength')} | Ловкость (Л): {c.get('Agility')} | Мастерство (М): {c.get('Mastery')} | Смекалка (С): {c.get('Wit')} *(Сумма: {tot})*")
            if str(c.get('Description', '')).strip():
                md.append(f"- **Эффект / Описание:** {c.get('Description')}")
            md.append("")

    if "Фракции и Организации" in sections and not factions_df.empty:
        md.append("## 🏛 3. Фракции, Кланы и Организации\n")
        for _, f in factions_df.iterrows():
            cat_str = f.get('Org_Category', 'Группировка')
            md.append(f"### 🏛 {f.get('Name')} ({cat_str})")
            if f.get('Motto'): md.append(f"*{f.get('Motto')}*")
            md.append(f"- **Лидер:** {f.get('Leader')} | **Идеология:** {f.get('Ideology')} | **Штаб:** {f.get('Base_Location')} | **Влияние:** {f.get('Influence')}/100")
            if f.get('Allies'): md.append(f"- **Союзники:** {f.get('Allies')}")
            if f.get('Enemies'): md.append(f"- **Враги:** {f.get('Enemies')}")
            if f.get('Description'): md.append(f"\n{f.get('Description')}")
            md.append("")

    if "Карта Местности (Локации)" in sections and not locations_df.empty:
        md.append("## 🗺 4. Локации и География\n")
        for _, l in locations_df.iterrows():
            md.append(f"### 📍 {l.get('Name')} ({l.get('Region')} — {l.get('Loc_Type')})")
            md.append(f"- **Контроль:** {l.get('Controlling_Faction')} | **Координаты:** X={l.get('X')}, Y={l.get('Y')}")
            if l.get('Linked_Characters'): md.append(f"- **Связанные герои:** {l.get('Linked_Characters')}")
            if l.get('Description'): md.append(f"\n{l.get('Description')}")
            md.append("")

    if "Хронология Вселенной" in sections and not timeline_df.empty:
        md.append("## ⏳ 5. Хронология Вселенной\n")
        t_sorted = timeline_df.sort_values(by='Order_Index')
        for _, t in t_sorted.iterrows():
            md.append(f"### [{t.get('Order_Index')}] {t.get('Date_Label')} ({t.get('Era')}) — {t.get('Title')}")
            if t.get('Location'): md.append(f"- **Место:** {t.get('Location')}")
            if t.get('Participants'): md.append(f"- **Участники:** {t.get('Participants')}")
            if t.get('Description'): md.append(f"\n{t.get('Description')}")
            md.append("")

    if "Архивы и Лор" in sections and not archives_df.empty:
        md.append("## 📚 6. Архивы и Лор\n")
        for _, a in archives_df.iterrows():
            md.append(f"### 📄 {a.get('Title')} [{a.get('Category')}]")
            md.append(str(a.get('Content', '')) + "\n")

    if "Сценарии и Сюжет" in sections and not scripts_df.empty:
        md.append("## 🎬 7. Сценарии и Сюжетные Линии\n")
        for _, s in scripts_df.iterrows():
            md.append(f"### 🎭 {s.get('Title')} (Арка: {s.get('Act')} | Статус: {s.get('Status')})")
            md.append(str(s.get('Content', '')) + "\n")

    return "\n".join(md)

def generate_universe_docx(md_text):
    doc = docx.Document()
    for line in md_text.split('\n'):
        stripped = line.strip()
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=0)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith('- '):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        elif stripped == "":
            continue
        else:
            doc.add_paragraph(line)
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()

# --- ИНИЦИАЛИЗАЦИЯ ДАННЫХ ---
if 'nodes' not in st.session_state:
    n, e, a, c, s, loc, tm, fac = load_all_data()
    st.session_state['nodes'] = n
    st.session_state['edges'] = e
    st.session_state['archives'] = a
    st.session_state['cards'] = c
    st.session_state['scripts'] = s
    st.session_state['locations'] = loc
    st.session_state['timeline'] = tm
    st.session_state['factions'] = fac

if 'config' not in st.session_state:
    st.session_state['config'] = load_config()
if 'colors' not in st.session_state['config']: st.session_state['config']['colors'] = {}
if 'options' not in st.session_state['config']: st.session_state['config']['options'] = {}
if 'region_maps' not in st.session_state['config']: st.session_state['config']['region_maps'] = {}

# --- ЕДИНЫЙ ИНДЕКС ПЕРСОНАЖЕЙ ДЛЯ ПОДСКАЗОК ВО ВСЕХ МОДУЛЯХ ---
CHAR_LOOKUP = build_universe_character_lookup(
    st.session_state['nodes'],
    st.session_state['edges'],
    st.session_state['cards']
)
CHAR_MATCHERS = build_character_matchers(list(CHAR_LOOKUP.keys()))
CHAR_POPUP_SPANS = {name: build_inline_popup_span(name, info) for name, info in CHAR_LOOKUP.items()}

# --- БОКОВОЕ МЕНЮ НАВИГАЦИИ ---
st.sidebar.title("🌆 БИГ СИТИ")
app_mode = st.sidebar.radio(
    "Навигация по экосистеме:",
    [
        "🌐 Персонажи и Граф",
        "🃏 Галерея Карточек",
        "🏛 Фракции и Организации",
        "🗺 Карта Местности",
        "⏳ Хронология Вселенной",
        "📚 Архивы и Лор",
        "🎬 Сценарии и Сюжет",
        "📥 Экспорт «Библии Вселенной»"
    ]
)

st.sidebar.write("---")
if st.sidebar.button("🔄 Синхронизировать с облаком", use_container_width=True):
    st.cache_data.clear()
    n, e, a, c, s, loc, tm, fac = load_all_data()
    st.session_state['nodes'], st.session_state['edges'] = n, e
    st.session_state['archives'], st.session_state['cards'], st.session_state['scripts'] = a, c, s
    st.session_state['locations'], st.session_state['timeline'], st.session_state['factions'] = loc, tm, fac
    st.session_state['config'] = load_config()
    st.rerun()

# Глобальная быстрая справка по любому персонажу прямо в боковой панели
if CHAR_LOOKUP:
    with st.sidebar.expander("🔍 Быстрая справка по герою", expanded=False):
        quick_hero = st.selectbox("Выберите персонажа:", sorted(list(CHAR_LOOKUP.keys())), key="sb_quick_hero")
        if quick_hero:
            q_info = CHAR_LOOKUP[quick_hero]
            if q_info.get("has_card"):
                st.markdown(f"**{q_info['card_name']}** ({q_info['alignment']} • {q_info['faction']})")
                if q_info.get("image"):
                    st.image(q_info["image"], use_container_width=True)
                st.code(f"С: {q_info['str']} | Л: {q_info['agi']}\nМ: {q_info['mas']} | С: {q_info['wit']}")
                if q_info.get("desc"):
                    st.caption(q_info["desc"])
            if q_info.get("lore"):
                for l_line in q_info["lore"]:
                    st.caption(f"▪ {l_line}")
            if q_info.get("relations"):
                st.caption("🔗 **Связи:** " + "; ".join(q_info["relations"]))


# ==============================================================================
# МОДУЛЬ 1: ПЕРСОНАЖИ, ГРАФ, ДОСЬЕ (С ПЕРЕКРЕСТНЫМИ ССЫЛКАМИ) И СТАТИСТИКА
# ==============================================================================
if app_mode == "🌐 Персонажи и Граф":
    st.title("🌐 База Персонажей и Карта Связей")
    tab_graph, tab_dossier, tab_chars, tab_edges, tab_stats = st.tabs([
        "🕸 Интерактивный Граф", "🗂 Досье и Перекрестные ссылки", "👤 Таблица Персонажей", "🔗 Таблица Связей", "📊 Статистика"
    ])

    with tab_dossier:
        all_chars_list = st.session_state['nodes']['Имя / Название персонажа'].dropna().unique().tolist()
        if not all_chars_list:
            st.info("База персонажей пока пуста.")
        else:
            sel_hero = st.selectbox("🔍 Выберите персонажа для просмотра полного досье:", all_chars_list)
            hero_row = st.session_state['nodes'][st.session_state['nodes']['Имя / Название персонажа'] == sel_hero].iloc[0]
            
            d_col1, d_col2 = st.columns([1.2, 1.8])
            with d_col1:
                st.subheader(f"👤 {sel_hero}")
                for col_n in st.session_state['nodes'].columns:
                    if col_n not in ['Имя / Название персонажа', 'Заметки']:
                        st.markdown(f"**{col_n}:** {hero_row.get(col_n, 'Не указано')}")
                if str(hero_row.get('Заметки', '')).strip() not in ['', 'Не указано']:
                    st.markdown("**Заметки / Лор:**")
                    hl_notes, _ = render_text_with_spoilers_and_characters(
                        str(hero_row['Заметки']), CHAR_MATCHERS, CHAR_POPUP_SPANS, enable_spoilers=False
                    )
                    st.markdown(f"<div class='script-body'>{hl_notes}</div>", unsafe_allow_html=True)
                
                st.write("---")
                st.markdown("#### 🔗 Связи на графе")
                hero_rels = get_character_relations(sel_hero, st.session_state['edges'])
                if hero_rels:
                    for r in hero_rels:
                        t_low = r['type'].lower()
                        css_c = "rel-enemy" if t_low in ['вражда', 'конфликт'] else ("rel-ally" if t_low in ['союз', 'родство', 'подчинение'] else "rel-neutral")
                        st.markdown(f"""
                        <div class="rel-pill {css_c}">
                            <span><b>{html.escape(r['partner'])}</b> ({html.escape(r['type'])})</span>
                            <span>Сила: {'⭐' * r['weight']}</span>
                        </div>
                        """, unsafe_allow_html=True)
                else:
                    st.caption("У этого персонажа пока нет зафиксированных связей.")

            with d_col2:
                st.subheader("🃏 Привязанные карточки и бонусы")
                linked_cards = get_cards_for_character(sel_hero, st.session_state['cards'])
                if not linked_cards.empty:
                    c_cols = st.columns(2)
                    for idx_c, (_, crd) in enumerate(linked_cards.iterrows()):
                        with c_cols[idx_c % 2]:
                            al_str = str(crd['Alignment'])
                            al_cls = "align-hero" if al_str == "Герой" else ("align-villain" if al_str == "Злодей" else "align-neutral")
                            st.markdown(f"""
                            <div class="game-card">
                                <div class="game-card-header">
                                    <span style="color:#ffcc00; font-size:12px;">{html.escape(str(crd['Faction']))}</span>
                                    <span class="game-card-type">{html.escape(str(crd['Card_Type']))}</span>
                                </div>
                                <div class="game-card-title">{html.escape(str(crd['Name']))}</div>
                            </div>
                            """, unsafe_allow_html=True)
                            if crd['Image_Data']:
                                st.image(crd['Image_Data'], use_container_width=True)
                            st.markdown(f"""
                            <div class="stats-grid">
                                <div class="stat-item"><span class="stat-letter">С</span>{int(crd['Strength'])}</div>
                                <div class="stat-item"><span class="stat-letter">Л</span>{int(crd['Agility'])}</div>
                                <div class="stat-item"><span class="stat-letter">М</span>{int(crd['Mastery'])}</div>
                                <div class="stat-item"><span class="stat-letter">С</span>{int(crd['Wit'])}</div>
                            </div>
                            <div class="alignment-badge {al_cls}">{html.escape(al_str)}</div>
                            """, unsafe_allow_html=True)
                            if crd['Description']:
                                st.caption(crd['Description'])
                else:
                    st.info("К этому персонажу пока не привязано ни одной игровой карточки в Галерее.")

            # --- БЛОК ПЕРЕКРЕСТНЫХ ССЫЛОК ПО ВСЕЙ ЭКОСИСТЕМЕ ---
            st.write("---")
            st.subheader("📖 Перекрестные ссылки в экосистеме (Авто-поиск упоминаний)")
            cross_refs = find_character_mentions_across_ecosystem(
                sel_hero, CHAR_MATCHERS,
                st.session_state['archives'],
                st.session_state['scripts'],
                st.session_state['locations'],
                st.session_state['timeline']
            )
            cr1, cr2, cr3, cr4 = st.columns(4)
            with cr1:
                st.markdown(f"**📚 Архивы ({len(cross_refs['archives'])})**")
                if cross_refs['archives']:
                    for item in cross_refs['archives']:
                        st.markdown(f"▪ **{item['title']}** *({item['cat']})*")
                        if item['snippet']:
                            st.caption(f"«{item['snippet']}»")
                else:
                    st.caption("Не упоминается в архивах.")
            with cr2:
                st.markdown(f"**🎬 Сценарии ({len(cross_refs['scripts'])})**")
                if cross_refs['scripts']:
                    for item in cross_refs['scripts']:
                        st.markdown(f"▪ **{item['title']}** *({item['act']} • {item['status']})*")
                        if item['snippet']:
                            st.caption(f"«{item['snippet']}»")
                else:
                    st.caption("Не упоминается в сценариях.")
            with cr3:
                st.markdown(f"**🗺 Локации ({len(cross_refs['locations'])})**")
                if cross_refs['locations']:
                    for item in cross_refs['locations']:
                        st.markdown(f"▪ 📍 **{item['name']}** *({item['region']} • {item['type']})*")
                else:
                    st.caption("Не привязан к локациям.")
            with cr4:
                st.markdown(f"**⏳ Хронология ({len(cross_refs['timeline'])})**")
                if cross_refs['timeline']:
                    for item in cross_refs['timeline']:
                        st.markdown(f"▪ 🕰 **{item['date']}**: {item['title']} *({item['era']})*")
                else:
                    st.caption("Нет событий в хронологии.")

    with tab_stats:
        st.subheader("Аналитика и распределение сил")
        m1, m2, m3 = st.columns(3)
        m1.metric("👥 Всего персонажей", len(st.session_state['nodes']))
        m2.metric("🔗 Всего связей", len(st.session_state['edges']))
        color_cols = [c for c in st.session_state['nodes'].columns if c not in ["Имя / Название персонажа", "Заметки"]]
        m3.metric("🗂 Создано категорий", len(color_cols))
        st.write("---")
        
        c_stat1, c_stat2 = st.columns(2)
        with c_stat1:
            st.markdown("**Распределение по категориям**")
            stat_cat = st.selectbox("Выберите категорию:", color_cols)
            if stat_cat and not st.session_state['nodes'].empty:
                counts = st.session_state['nodes'][stat_cat].value_counts().reset_index()
                counts.columns = [stat_cat, 'Количество']
                fig_pie = px.pie(counts, values='Количество', names=stat_cat, hole=0.4, color_discrete_sequence=px.colors.qualitative.Pastel)
                fig_pie.update_traces(textposition='inside', textinfo='percent+label')
                fig_pie.update_layout(showlegend=False, margin=dict(t=10, b=10, l=10, r=10), paper_bgcolor='rgba(0,0,0,0)', font_color='white')
                st.plotly_chart(fig_pie, use_container_width=True)
                
        with c_stat2:
            st.markdown("**👑 Топ-10 Влиятельных фигур**")
            if not st.session_state['edges'].empty:
                all_edges = pd.concat([st.session_state['edges']['Source'], st.session_state['edges']['Target']])
                top_chars = all_edges.value_counts().head(10).reset_index()
                top_chars.columns = ['Персонаж', 'Количество связей']
                fig_bar = px.bar(top_chars, x='Количество связей', y='Персонаж', orientation='h', color='Количество связей', color_continuous_scale='Blues')
                fig_bar.update_layout(yaxis={'categoryorder':'total ascending'}, coloraxis_showscale=False, margin=dict(t=10, b=10, l=10, r=10), paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font_color='white')
                st.plotly_chart(fig_bar, use_container_width=True)
            else:
                st.info("Добавьте связи, чтобы увидеть лидеров.")

    with tab_chars:
        protected_cols = ["Имя / Название персонажа", "Заметки"]
        custom_cols = [c for c in st.session_state['nodes'].columns if c not in protected_cols]

        with st.expander("🔄 Умное переименование персонажа (С сохранением связей и карточек)"):
            all_chars = st.session_state['nodes']['Имя / Название персонажа'].dropna().unique().tolist()
            old_name = st.selectbox("Выберите персонажа:", all_chars) if all_chars else None
            new_char_name = st.text_input("Новое имя:").strip()
            if st.button("Применить новое имя") and old_name:
                if new_char_name and new_char_name != old_name and new_char_name not in all_chars:
                    st.session_state['nodes']['Имя / Название персонажа'] = st.session_state['nodes']['Имя / Название персонажа'].replace(old_name, new_char_name)
                    st.session_state['edges']['Source'] = st.session_state['edges']['Source'].replace(old_name, new_char_name)
                    st.session_state['edges']['Target'] = st.session_state['edges']['Target'].replace(old_name, new_char_name)
                    st.session_state['cards']['Linked_Character'] = st.session_state['cards']['Linked_Character'].replace(old_name, new_char_name)
                    save_sheet("Nodes", st.session_state['nodes'])
                    save_sheet("Edges", st.session_state['edges'])
                    save_sheet("Cards", st.session_state['cards'][CARD_COLS])
                    st.success("Имя обновлено в базе, на графе и в карточках!")
                    st.rerun()

        col_menu1, col_menu2 = st.columns(2)
        with col_menu1:
            with st.expander("⚙️ Управление столбцами"):
                new_col = st.text_input("Новый столбец (напр. Клан, Оружие, Раса):").strip()
                if st.button("➕ Создать столбец"):
                    if new_col and new_col not in st.session_state['nodes'].columns:
                        st.session_state['nodes'][new_col] = "Не указано"
                        save_sheet("Nodes", st.session_state['nodes'])
                        st.rerun()
                if custom_cols:
                    edit_col = st.selectbox("Столбец:", custom_cols)
                    c_ren, c_del = st.columns(2)
                    with c_ren:
                        new_name = st.text_input("Новое название:").strip()
                        if st.button("✏️ Переименовать") and new_name:
                            st.session_state['nodes'] = st.session_state['nodes'].rename(columns={edit_col: new_name})
                            save_sheet("Nodes", st.session_state['nodes'])
                            st.rerun()
                    with c_del:
                        st.write("")
                        if st.button("🗑 Удалить"):
                            st.session_state['nodes'] = st.session_state['nodes'].drop(columns=[edit_col])
                            save_sheet("Nodes", st.session_state['nodes'])
                            st.rerun()

        with col_menu2:
            with st.expander("🎨 Настройка вариантов и цветов"):
                if custom_cols:
                    color_col = st.selectbox("Категория:", custom_cols, key="color_col")
                    with st.form("add_variant_form"):
                        nv_name = st.text_input("Новый класс/вариант:").strip()
                        nv_color = st.color_picker("Цвет на графе:", "#4da6ff")
                        if st.form_submit_button("Сохранить вариант и цвет"):
                            if nv_name:
                                st.session_state['config']['colors'].setdefault(color_col, {})[nv_name] = nv_color
                                opts = st.session_state['config']['options'].setdefault(color_col, [])
                                if nv_name not in opts: opts.append(nv_name)
                                save_config_to_cloud(st.session_state['config'])
                                st.success(f"'{nv_name}' добавлен!")
                                st.rerun()

        # Синхронизация выпадающих списков таблицы Графа с модулем Фракций, Кланов и Организаций
        col_config = {
            "Имя / Название персонажа": st.column_config.TextColumn("Имя (уникальное)", required=True),
            "Заметки": st.column_config.TextColumn("Способности / Лор")
        }
        
        df_fac_sync = st.session_state.get('factions', pd.DataFrame(columns=FAC_COLS))
        all_fac_names = df_fac_sync['Name'].dropna().unique().tolist() if not df_fac_sync.empty else []

        for col in custom_cols:
            existing_opts = [str(x) for x in st.session_state['nodes'][col].dropna().unique() if str(x).strip()]
            saved_opts = st.session_state['config']['options'].get(col, [])
            
            matched_facs = []
            if not df_fac_sync.empty and 'Org_Category' in df_fac_sync.columns:
                matched_facs = df_fac_sync[
                    df_fac_sync['Org_Category'].astype(str).str.strip().str.lower() == col.strip().lower()
                ]['Name'].dropna().tolist()
            
            if col.strip().lower() in ["группировка", "фракция", "организация", "клан"]:
                matched_facs = list(set(matched_facs + all_fac_names))

            all_opts = sorted(list(set(existing_opts + saved_opts + matched_facs + ["Не указано"])))
            col_config[col] = st.column_config.SelectboxColumn(col, options=all_opts)

        edited_nodes = st.data_editor(st.session_state['nodes'], num_rows="dynamic", use_container_width=True, height=550, column_config=col_config)
        if st.button("💾 Сохранить персонажей", type="primary"):
            df_to_save = edited_nodes.dropna(subset=['Имя / Название персонажа']).fillna("Не указано")
            st.session_state['nodes'] = df_to_save
            save_sheet("Nodes", df_to_save)
            st.success("Персонажи сохранены!")
            st.rerun()

    with tab_edges:
        valid_chars = st.session_state['nodes']['Имя / Название персонажа'].dropna().unique().tolist()
        edited_edges = st.data_editor(
            st.session_state['edges'], num_rows="dynamic", use_container_width=True, height=500,
            column_config={
                "Source": st.column_config.SelectboxColumn("Кто (Source)", options=valid_chars, required=True),
                "Target": st.column_config.SelectboxColumn("С кем (Target)", options=valid_chars, required=True),
                "Type": st.column_config.SelectboxColumn("Тип связи", options=["Союз", "Вражда", "Родство", "Подчинение", "Нейтрально", "Конфликт"]),
                "Weight": st.column_config.NumberColumn("Сила связи (1-5)", min_value=1, max_value=5, step=1)
            }
        )
        if st.button("💾 Сохранить связи", type="primary"):
            df_edges_save = edited_edges.dropna(subset=['Source', 'Target'])
            st.session_state['edges'] = df_edges_save
            save_sheet("Edges", df_edges_save)
            st.success("Связи сохранены!")
            st.rerun()

    with tab_graph:
        color_options = [c for c in st.session_state['nodes'].columns if c not in ["Имя / Название персонажа", "Заметки"]]
        if color_options:
            color_by = st.sidebar.selectbox("🎨 Раскрашивать граф по:", color_options)
            show_card_avatars = st.sidebar.checkbox("🖼 Арт карточки внутри кружочка узла", value=True)
            
            active_filters = {}
            with st.sidebar.expander("🌪 Фильтры графа", expanded=False):
                for col in color_options:
                    unique_vals = st.session_state['nodes'][col].dropna().unique().tolist()
                    active_filters[col] = st.multiselect(f"{col}:", unique_vals, default=unique_vals)

            df_n = st.session_state['nodes'].copy()
            for col, selected_vals in active_filters.items():
                df_n = df_n[df_n[col].isin(selected_vals)]
                
            valid_names = df_n['Имя / Название персонажа'].tolist()
            df_e = st.session_state['edges']
            filtered_edges = df_e[(df_e['Source'].isin(valid_names)) & (df_e['Target'].isin(valid_names))]

            net = Network(height='720px', width='100%', bgcolor='#181820', font_color='white')
            net.force_atlas_2based(gravity=-35)

            unique_color_targets = df_n[color_by].dropna().unique().tolist()
            node_popup_dict = {}

            for _, row in df_n.iterrows():
                name = str(row['Имя / Название персонажа']).strip()
                plan = str(row.get('Категория плана', '3 план'))
                val = str(row.get(color_by, 'Не указано'))
                
                saved_color = st.session_state['config']['colors'].get(color_by, {}).get(val)
                if not saved_color:
                    color_idx = unique_color_targets.index(val) if val in unique_color_targets else 0
                    saved_color = PALETTE[color_idx % len(PALETTE)]
                    if val == "Не указано":
                        saved_color = "#ffffff"
                        
                size = 28 if '1' in plan else (18 if '2' in plan else 12)
                info = CHAR_LOOKUP.get(name, {"has_card": False, "lore": []})
                node_popup_dict[name] = info

                if info.get("has_card") and show_card_avatars and info.get("image"):
                    net.add_node(
                        name,
                        label=name,
                        size=size + 6,
                        shape='circularImage',
                        image=info["image"],
                        borderWidth=4,
                        color={'border': saved_color, 'background': '#1b1b26', 'highlight': {'border': '#00d4ff'}}
                    )
                else:
                    net.add_node(name, label=name, size=size, color=saved_color)

            added_nodes = net.get_nodes()
            for _, row in filtered_edges.iterrows():
                src, tgt = str(row['Source']).strip(), str(row['Target']).strip()
                rel_type = str(row.get('Type', 'Связь')).strip()
                try:
                    weight = float(row.get('Weight', 1.0))
                except ValueError:
                    weight = 1.0
                if src in added_nodes and tgt in added_nodes:
                    edge_color = '#ff4d4d' if rel_type.lower() in ['вражда', 'конфликт'] else ('#4da6ff' if rel_type.lower() in ['союз', 'родство', 'подчинение'] else '#888888')
                    net.add_edge(src, tgt, value=weight, color=edge_color, title=rel_type)

            net.save_graph('pyvis_graph.html')
            with open('pyvis_graph.html', 'r', encoding='utf-8') as f:
                raw_graph_html = f.read()
            enhanced_graph_html = inject_card_hover_into_graph(raw_graph_html, node_popup_dict)
            components.html(enhanced_graph_html, height=750)


# ==============================================================================
# МОДУЛЬ 2: ГАЛЕРЕЯ КАРТОЧЕК, ИНДИКАТОР БАЛАНСА И ПРЯМАЯ ПОКАЗАТЕЛЕЙ (ВСЕ СТАТЫ)
# ==============================================================================
elif app_mode == "🃏 Галерея Карточек":
    st.title("🃏 Игровые Карточки БИГ СИТИ")
    df_cards = st.session_state['cards'].copy()
    for s_col in ['Strength', 'Agility', 'Mastery', 'Wit']:
        df_cards[s_col] = pd.to_numeric(df_cards[s_col], errors='coerce').fillna(0).astype(int)
    df_cards['Total_Power'] = df_cards['Strength'] + df_cards['Agility'] + df_cards['Mastery'] + df_cards['Wit']

    valid_graph_chars = st.session_state['nodes']['Имя / Название персонажа'].dropna().astype(str).unique().tolist()
    tab_gallery, tab_number_lines, tab_card_stats = st.tabs([
        "🃏 Коллекция и Конструктор",
        "📏 Прямая Показателей (Все статы)",
        "📊 Статистика и Баланс Карточек"
    ])

    with tab_gallery:
        with st.expander("➕ Добавить или отредактировать карточку", expanded=True):
            edit_mode = st.selectbox("Выберите карточку для редактирования или создайте новую:", ["Создать новую карточку"] + df_cards['Name'].tolist())
            
            if edit_mode == "Создать новую карточку":
                c_name_val, c_type_val, c_align_val, c_fac_val, c_desc_val, c_img_val, c_link_val = "", "Персонаж", "Нейтралитет", "Общие", "", "", "— Авто / По имени —"
                s_str, s_agi, s_mas, s_wit = 8000, 8000, 8000, 8000
            else:
                row_c = df_cards[df_cards['Name'] == edit_mode].iloc[0]
                c_name_val = str(row_c['Name'])
                c_type_val = str(row_c['Card_Type']) if row_c['Card_Type'] in ["Персонаж", "Бонус"] else "Персонаж"
                c_align_val = str(row_c['Alignment']) if row_c['Alignment'] in ["Герой", "Злодей", "Нейтралитет"] else "Нейтралитет"
                c_fac_val = str(row_c['Faction']) or "Общие"
                c_desc_val = str(row_c['Description'])
                c_img_val = str(row_c['Image_Data'])
                saved_link = str(row_c.get('Linked_Character', '')).strip()
                c_link_val = saved_link if saved_link in valid_graph_chars else "— Авто / По имени —"
                s_str, s_agi, s_mas, s_wit = int(row_c['Strength']), int(row_c['Agility']), int(row_c['Mastery']), int(row_c['Wit'])

            col_f1, col_f2 = st.columns(2)
            with col_f1:
                c_name = st.text_input("Имя / Название карточки (напр. ТЕР РОТ):", value=c_name_val)
                tc1, tc2 = st.columns(2)
                with tc1:
                    type_opts = ["Персонаж", "Бонус"]
                    c_type = st.selectbox("Категория карточки:", type_opts, index=type_opts.index(c_type_val))
                with tc2:
                    align_opts = ["Герой", "Злодей", "Нейтралитет"]
                    c_align = st.selectbox("Класс (снизу карточки):", align_opts, index=align_opts.index(c_align_val))
                
                link_options = ["— Авто / По имени —"] + valid_graph_chars
                c_link_idx = link_options.index(c_link_val) if c_link_val in link_options else 0
                c_linked_char = st.selectbox("🔗 Привязка к профилю на графе:", link_options, index=c_link_idx)
                
                # Выпадающий список фракций/кланов, синхронизированный с модулем Фракций
                fac_from_mod = st.session_state['factions']['Name'].dropna().unique().tolist() if not st.session_state['factions'].empty else []
                fac_from_cards = [f for f in df_cards['Faction'].dropna().unique().tolist() if str(f).strip()]
                all_known_facs = sorted(list(set(["Общие", "Нейтралитет"] + fac_from_mod + fac_from_cards)))
                if c_fac_val and c_fac_val not in all_known_facs:
                    all_known_facs.append(c_fac_val)

                sel_fac_card = st.selectbox(
                    "🏛 Фракция / Клан / Организация:",
                    all_known_facs + ["➕ Ввести новую вручную..."],
                    index=all_known_facs.index(c_fac_val) if c_fac_val in all_known_facs else 0
                )
                if sel_fac_card == "➕ Ввести новую вручную...":
                    c_fac = st.text_input("Введите название новой фракции/клана:", value="")
                else:
                    c_fac = sel_fac_card

            with col_f2:
                st.markdown("**Показатели карточки (или прибавка от бонуса):**")
                st1, st2 = st.columns(2)
                with st1:
                    val_str = st.number_input("💪 Сила (С):", value=s_str, step=10)
                    val_mas = st.number_input("🛠 Мастерство (М):", value=s_mas, step=10)
                with st2:
                    val_agi = st.number_input("⚡ Ловкость (Л):", value=s_agi, step=10)
                    val_wit = st.number_input("🧠 Смекалка (С):", value=s_wit, step=10)
                
                val_total = int(val_str) + int(val_agi) + int(val_mas) + int(val_wit)
                st.caption(f"⚡ **Суммарная мощь текущего ввода:** `{val_total}`")
                auto_add_to_graph = st.checkbox(
                    "🌐 Автоматически добавить персонажа в базу Графа (если его там ещё нет)",
                    value=(c_type == "Персонаж")
                )

            # ==================================================================
            # ⚖️ ЖИВОЙ БАЛАНС И ПРЯМАЯ ПОКАЗАТЕЛЕЙ ПРЯМО ПОД ПОЛЯМИ ВВОДА
            # ==================================================================
            st.markdown('<div class="balance-box">', unsafe_allow_html=True)
            st.markdown("#### ⚖️ Живой баланс и прямая показателей")
            
            other_cards_df = df_cards[df_cards['Name'] != edit_mode].copy() if edit_mode != "Создать новую карточку" else df_cards.copy()
            other_pers_df = other_cards_df[other_cards_df['Card_Type'] == c_type].copy()
            if other_pers_df.empty:
                other_pers_df = other_cards_df.copy()

            b_ctrl1, b_ctrl2 = st.columns([1.3, 1.7])
            with b_ctrl1:
                axis_stat_choice = st.selectbox(
                    "📏 Показать числовую прямую для характеристики:",
                    ["💪 Сила (С)", "⚡ Ловкость (Л)", "🛠 Мастерство (М)", "🧠 Смекалка (С)", "🔥 Суммарная мощь (С+Л+М+С)"],
                    key="live_axis_stat"
                )
            with b_ctrl2:
                tolerance_val = st.slider(
                    "🎯 Допустимый разброс (± единиц) для поиска ровни и подсветки коридора:",
                    min_value=0, max_value=5000, value=500, step=50, key="live_tol_slider"
                )

            stat_map_live = {
                "💪 Сила (С)": ("Strength", int(val_str)),
                "⚡ Ловкость (Л)": ("Agility", int(val_agi)),
                "🛠 Мастерство (М)": ("Mastery", int(val_mas)),
                "🧠 Смекалка (С)": ("Wit", int(val_wit)),
                "🔥 Суммарная мощь (С+Л+М+С)": ("Total_Power", int(val_total))
            }
            chosen_col_name, chosen_curr_val = stat_map_live[axis_stat_choice]
            preview_label = c_name.strip() if c_name.strip() else "Вводимая карточка"

            fig_live_line = build_stat_number_line_figure(
                other_pers_df,
                chosen_col_name,
                axis_stat_choice,
                highlight_val=chosen_curr_val,
                highlight_name=preview_label,
                tolerance=tolerance_val if chosen_col_name != "Total_Power" else tolerance_val * 2,
                height=210
            )
            st.plotly_chart(fig_live_line, use_container_width=True)

            st.markdown(f"**📋 Сравнение с существующими карточками в диапазоне ±{tolerance_val} (и ближайшие соседи):**")
            stat_check_list = [
                ("💪 Сила (С)", "Strength", int(val_str), tolerance_val),
                ("⚡ Ловкость (Л)", "Agility", int(val_agi), tolerance_val),
                ("🛠 Мастерство (М)", "Mastery", int(val_mas), tolerance_val),
                ("🧠 Смекалка (С)", "Wit", int(val_wit), tolerance_val),
                ("🔥 Сумма", "Total_Power", int(val_total), tolerance_val * 2)
            ]
            s_cols = st.columns(5)
            for idx_s, (s_title, s_col_key, curr_v, tol_v) in enumerate(stat_check_list):
                with s_cols[idx_s]:
                    st.markdown(f"**{s_title}: `{curr_v}`**")
                    if other_pers_df.empty:
                        st.caption("База пуста")
                    else:
                        other_pers_df['diff'] = other_pers_df[s_col_key] - curr_v
                        in_range = other_pers_df[other_pers_df['diff'].abs() <= tol_v].sort_values(by='diff', ascending=False)
                        
                        weaker_df = other_pers_df[other_pers_df[s_col_key] < curr_v].sort_values(by=s_col_key, ascending=False)
                        stronger_df = other_pers_df[other_pers_df[s_col_key] > curr_v].sort_values(by=s_col_key, ascending=True)

                        if not stronger_df.empty:
                            up_r = stronger_df.iloc[0]
                            st.caption(f"🔼 Сильнее: **{up_r['Name']}** ({up_r[s_col_key]}, `+{up_r[s_col_key]-curr_v}`)")
                        else:
                            st.caption("🔼 Сильнее: *нет (Топ-1)*")

                        if not weaker_df.empty:
                            dn_r = weaker_df.iloc[0]
                            st.caption(f"🔽 Слабее: **{dn_r['Name']}** ({dn_r[s_col_key]}, `{dn_r[s_col_key]-curr_v}`)")
                        else:
                            st.caption("🔽 Слабее: *нет (Минимум)*")

                        st.markdown(f"<div style='font-size:12px; color:#00d4ff; margin-top:4px;'>В коридоре ±{tol_v} ({len(in_range)}):</div>", unsafe_allow_html=True)
                        if not in_range.empty:
                            for _, ir in in_range.head(6).iterrows():
                                d_int = int(ir['diff'])
                                d_str = f"+{d_int}" if d_int > 0 else ("0" if d_int == 0 else str(d_int))
                                st.markdown(
                                    f"<div class='neighbor-card'><span><b>{html.escape(str(ir['Name']))}</b></span>"
                                    f"<span>{int(ir[s_col_key])} (<b>{d_str}</b>)</span></div>",
                                    unsafe_allow_html=True
                                )
                        else:
                            st.caption("Никого в диапазоне.")
            st.markdown('</div>', unsafe_allow_html=True)

            img_col1, img_col2 = st.columns(2)
            with img_col1:
                uploaded_img = st.file_uploader("📸 Загрузить фото карточки (с телефона или ПК):", type=["png", "jpg", "jpeg", "webp"])
            with img_col2:
                c_url = st.text_input("ИЛИ прямая ссылка на изображение:", value="" if c_img_val.startswith("data:image") else c_img_val)

            c_desc = st.text_area("Описание бонуса / Особое свойство / Лор:", value=c_desc_val, height=80)

            btn_c1, btn_c2 = st.columns([3, 1])
            with btn_c1:
                if st.button("💾 Сохранить карточку в базу", type="primary", use_container_width=True):
                    if c_name.strip():
                        final_img = c_img_val
                        if uploaded_img is not None:
                            final_img = compress_uploaded_image(uploaded_img)
                        elif c_url.strip():
                            final_img = c_url.strip()

                        resolved_link = c_name.strip() if c_linked_char == "— Авто / По имени —" else c_linked_char
                        clean_fac_val = c_fac.strip() or "Общие"

                        new_card_data = {
                            "Name": c_name.strip(),
                            "Card_Type": c_type,
                            "Alignment": c_align,
                            "Faction": clean_fac_val,
                            "Strength": int(val_str),
                            "Agility": int(val_agi),
                            "Mastery": int(val_mas),
                            "Wit": int(val_wit),
                            "Description": c_desc.strip(),
                            "Image_Data": final_img,
                            "Linked_Character": resolved_link
                        }

                        base_df = st.session_state['cards'][CARD_COLS]
                        if edit_mode == "Создать новую карточку":
                            st.session_state['cards'] = pd.concat([base_df, pd.DataFrame([new_card_data])], ignore_index=True)
                        else:
                            idx = base_df[base_df['Name'] == edit_mode].index[0]
                            for k, v in new_card_data.items():
                                st.session_state['cards'].at[idx, k] = v

                        save_sheet("Cards", st.session_state['cards'][CARD_COLS])

                        if auto_add_to_graph and c_type == "Персонаж" and resolved_link not in valid_graph_chars:
                            new_node_row = {col: "Не указано" for col in st.session_state['nodes'].columns}
                            new_node_row['Имя / Название персонажа'] = resolved_link
                            new_node_row['Категория плана'] = '1 план'
                            if 'Группировка' in new_node_row:
                                new_node_row['Группировка'] = clean_fac_val
                            if 'Амплуа' in new_node_row:
                                new_node_row['Амплуа'] = c_align
                            if 'Заметки' in new_node_row:
                                new_node_row['Заметки'] = c_desc.strip()
                            st.session_state['nodes'] = pd.concat([st.session_state['nodes'], pd.DataFrame([new_node_row])], ignore_index=True)
                            save_sheet("Nodes", st.session_state['nodes'])

                        st.success(f"Карточка «{c_name.strip()}» успешно сохранена и синхронизирована с графом!")
                        st.rerun()
                    else:
                        st.warning("Укажите название карточки!")
            with btn_c2:
                if edit_mode != "Создать новую карточку":
                    if st.button("🗑 Удалить карточку", use_container_width=True):
                        st.session_state['cards'] = st.session_state['cards'][st.session_state['cards']['Name'] != edit_mode].reset_index(drop=True)
                        save_sheet("Cards", st.session_state['cards'][CARD_COLS])
                        st.rerun()

        st.write("---")

        if not df_cards.empty:
            f1, f2, f3, f4 = st.columns(4)
            with f1:
                filter_type = st.selectbox("Тип карточки:", ["Все", "Персонаж", "Бонус"])
            with f2:
                filter_align = st.selectbox("Класс:", ["Все", "Герой", "Злодей", "Нейтралитет"])
            with f3:
                factions = ["Все"] + sorted([f for f in df_cards['Faction'].unique().tolist() if f])
                filter_fac = st.selectbox("Фракция:", factions)
            with f4:
                sort_by = st.selectbox("Сортировать по:", [
                    "По порядку добавления",
                    "Общая мощь (по убыванию)",
                    "Сила (по убыванию)",
                    "Ловкость (по убыванию)",
                    "Мастерство (по убыванию)",
                    "Смекалка (по убыванию)",
                    "По алфавиту (А-Я)"
                ])

            search_q = st.text_input("🔍 Быстрый поиск по имени карточки, персонажу или тексту эффекта:")

            view_cards = df_cards.copy()
            if filter_type != "Все":
                view_cards = view_cards[view_cards['Card_Type'] == filter_type]
            if filter_align != "Все":
                view_cards = view_cards[view_cards['Alignment'] == filter_align]
            if filter_fac != "Все":
                view_cards = view_cards[view_cards['Faction'] == filter_fac]
            if search_q.strip():
                q = search_q.strip().lower()
                view_cards = view_cards[
                    view_cards['Name'].str.lower().str.contains(q) |
                    view_cards['Description'].str.lower().str.contains(q) |
                    view_cards['Linked_Character'].astype(str).str.lower().str.contains(q)
                ]

            sort_map = {
                "Общая мощь (по убыванию)": ("Total_Power", False),
                "Сила (по убыванию)": ("Strength", False),
                "Ловкость (по убыванию)": ("Agility", False),
                "Мастерство (по убыванию)": ("Mastery", False),
                "Смекалка (по убыванию)": ("Wit", False),
                "По алфавиту (А-Я)": ("Name", True)
            }
            if sort_by in sort_map:
                col_s, asc = sort_map[sort_by]
                view_cards = view_cards.sort_values(by=col_s, ascending=asc)

            st.caption(f"Найдено карточек: **{len(view_cards)}**")
            cols = st.columns(3)
            for i, (_, card) in enumerate(view_cards.iterrows()):
                with cols[i % 3]:
                    align_str = str(card['Alignment'])
                    align_class = "align-hero" if align_str == "Герой" else ("align-villain" if align_str == "Злодей" else "align-neutral")
                    
                    st.markdown(f"""
                    <div class="game-card">
                        <div class="game-card-header">
                            <span style="color:#ffcc00; font-size:12px;">{html.escape(str(card['Faction']))}</span>
                            <span class="game-card-type">{html.escape(str(card['Card_Type']))}</span>
                        </div>
                        <div class="game-card-title">{html.escape(str(card['Name']))}</div>
                    </div>
                    """, unsafe_allow_html=True)
                    
                    if card['Image_Data']:
                        st.image(card['Image_Data'], use_container_width=True)
                        
                    st.markdown(f"""
                    <div class="stats-grid">
                        <div class="stat-item"><span class="stat-letter">С</span>{int(card['Strength'])}</div>
                        <div class="stat-item"><span class="stat-letter">Л</span>{int(card['Agility'])}</div>
                        <div class="stat-item"><span class="stat-letter">М</span>{int(card['Mastery'])}</div>
                        <div class="stat-item"><span class="stat-letter">С</span>{int(card['Wit'])}</div>
                    </div>
                    <div class="alignment-badge {align_class}">{html.escape(align_str)}</div>
                    """, unsafe_allow_html=True)

                    if card['Description']:
                        hl_desc, _ = render_text_with_spoilers_and_characters(
                            str(card['Description']), CHAR_MATCHERS, CHAR_POPUP_SPANS, enable_spoilers=False
                        )
                        st.markdown(f"<div style='font-size:13px; color:#cccccc; margin-bottom:6px;'>{hl_desc}</div>", unsafe_allow_html=True)

                    linked_hero = resolve_linked_character(card, valid_graph_chars)
                    if linked_hero:
                        rels = get_character_relations(linked_hero, st.session_state['edges'])
                        with st.expander(f"🔗 Связи на графе: {linked_hero} ({len(rels)})"):
                            hero_matches = st.session_state['nodes'][
                                st.session_state['nodes']['Имя / Название персонажа'].astype(str).str.strip() == str(linked_hero).strip()
                            ]
                            if not hero_matches.empty:
                                hero_info = hero_matches.iloc[0]
                                meta_parts = [f"**{c}**: {hero_info[c]}" for c in st.session_state['nodes'].columns if c not in ['Имя / Название персонажа', 'Заметки'] and str(hero_info[c]) != 'Не указано']
                                if meta_parts:
                                    st.markdown(" • ".join(meta_parts))
                                if str(hero_info.get('Заметки', '')).strip() not in ['', 'Не указано']:
                                    st.caption(f"📝 {hero_info['Заметки']}")

                            if rels:
                                for r in rels:
                                    t_low = r['type'].lower()
                                    css_c = "rel-enemy" if t_low in ['вражда', 'конфликт'] else ("rel-ally" if t_low in ['союз', 'родство', 'подчинение'] else "rel-neutral")
                                    icon = "⚔️" if css_c == "rel-enemy" else ("🤝" if css_c == "rel-ally" else "⚖️")
                                    st.markdown(f"""
                                    <div class="rel-pill {css_c}">
                                        <span>{icon} <b>{html.escape(r['partner'])}</b> — {html.escape(r['type'])}</span>
                                        <span>{'⭐' * r['weight']}</span>
                                    </div>
                                    """, unsafe_allow_html=True)
                            else:
                                st.caption("Связей на графе пока нет.")

                            st.markdown("---")
                            st.caption("➕ **Быстро добавить связь:**")
                            other_chars = [c for c in valid_graph_chars if c != linked_hero]
                            if other_chars:
                                qc1, qc2 = st.columns(2)
                                with qc1:
                                    q_target = st.selectbox("С кем:", other_chars, key=f"qt_{i}_{card['Name']}")
                                with qc2:
                                    q_type = st.selectbox("Тип:", ["Союз", "Вражда", "Родство", "Подчинение", "Конфликт", "Нейтрально"], key=f"qtp_{i}_{card['Name']}")
                                q_weight = st.slider("Сила связи:", 1, 5, 3, key=f"qw_{i}_{card['Name']}")
                                if st.button("Добавить связь на граф", key=f"qbtn_{i}_{card['Name']}", use_container_width=True):
                                    new_edge = pd.DataFrame([{"Source": linked_hero, "Target": q_target, "Type": q_type, "Weight": q_weight}])
                                    st.session_state['edges'] = pd.concat([st.session_state['edges'], new_edge], ignore_index=True)
                                    save_sheet("Edges", st.session_state['edges'])
                                    st.success("Связь добавлена!")
                                    st.rerun()
                    else:
                        st.caption("⚪ *Персонаж еще не привязан к графу*")

                    st.write("---")
        else:
            st.info("В базе пока нет карточек. Добавьте первую карточку через панель выше!")

    # ==========================================================================
    # ВКЛАДКА 2: 📏 ПРЯМАЯ ПОКАЗАТЕЛЕЙ (5 ПАРАЛЛЕЛЬНЫХ ОСЕЙ ВСЕХ СТАТОВ)
    # ==========================================================================
    with tab_number_lines:
        st.subheader("📏 Параллельные Числовые Прямые Вселенной (Анализ пустых зон и перекосов)")
        if df_cards.empty:
            st.info("Добавьте карточки в базу, чтобы увидеть распределение на числовых прямых.")
        else:
            nl_c1, nl_c2, nl_c3 = st.columns(3)
            with nl_c1:
                nl_type = st.selectbox("Категория карточек на осях:", ["Персонаж", "Бонус", "Все"], key="nl_type")
            with nl_c2:
                nl_facs = ["Все"] + sorted([f for f in df_cards['Faction'].unique().tolist() if f])
                nl_fac = st.selectbox("Фильтр по фракции:", nl_facs, key="nl_fac")
            with nl_c3:
                nl_highlight_card = st.selectbox("🎯 Подсветить карточку ромбом:", ["— Не выделять —"] + df_cards['Name'].tolist(), key="nl_hl")

            nl_df = df_cards.copy()
            if nl_type != "Все":
                nl_df = nl_df[nl_df['Card_Type'] == nl_type]
            if nl_fac != "Все":
                nl_df = nl_df[nl_df['Faction'] == nl_fac]

            hl_row = None
            if nl_highlight_card != "— Не выделять —":
                hl_matches = df_cards[df_cards['Name'] == nl_highlight_card]
                if not hl_matches.empty:
                    hl_row = hl_matches.iloc[0]

            five_axes = [
                ("Strength", "💪 Сила (С)"),
                ("Agility", "⚡ Ловкость (Л)"),
                ("Mastery", "🛠 Мастерство (М)"),
                ("Wit", "🧠 Смекалка (С)"),
                ("Total_Power", "🔥 Суммарная мощь (С+Л+М+С)")
            ]

            for col_k, label_k in five_axes:
                h_val = int(hl_row[col_k]) if hl_row is not None else None
                h_name = str(hl_row['Name']) if hl_row is not None else ""
                fig_axis = build_stat_number_line_figure(
                    nl_df, col_k, label_k,
                    highlight_val=h_val,
                    highlight_name=h_name,
                    tolerance=0,
                    height=185
                )
                st.plotly_chart(fig_axis, use_container_width=True)

    with tab_card_stats:
        if df_cards.empty:
            st.info("Добавьте карточки, чтобы увидеть статистику и баланс.")
        else:
            st.subheader("Баланс классов и статистика колоды")
            k1, k2, k3, k4 = st.columns(4)
            k1.metric("🃏 Всего карточек", len(df_cards))
            k2.metric("👤 Персонажей", len(df_cards[df_cards['Card_Type'] == 'Персонаж']))
            k3.metric("🎁 Карточек бонусов", len(df_cards[df_cards['Card_Type'] == 'Бонус']))
            char_only = df_cards[df_cards['Card_Type'] == 'Персонаж']
            avg_power = int(char_only['Total_Power'].mean()) if not char_only.empty else 0
            k4.metric("⚡ Средняя сумма статов", avg_power)

            st.write("---")
            ch1, ch2 = st.columns(2)
            with ch1:
                st.markdown("**Соотношение классов (Герой / Злодей / Нейтралитет)**")
                al_counts = df_cards['Alignment'].value_counts().reset_index()
                al_counts.columns = ['Класс', 'Количество']
                fig_al = px.pie(al_counts, values='Количество', names='Класс', hole=0.45, color='Класс', color_discrete_map=ALIGN_COLORS)
                fig_al.update_traces(textposition='inside', textinfo='percent+label')
                fig_al.update_layout(showlegend=False, margin=dict(t=10, b=10, l=10, r=10), paper_bgcolor='rgba(0,0,0,0)', font_color='white')
                st.plotly_chart(fig_al, use_container_width=True)

            with ch2:
                st.markdown("**Баланс средних показателей по классам (Персонажи)**")
                if not char_only.empty:
                    avg_stats = char_only.groupby('Alignment')[['Strength', 'Agility', 'Mastery', 'Wit']].mean().reset_index()
                    avg_melted = avg_stats.melt(id_vars='Alignment', var_name='Показатель', value_name='Среднее значение')
                    stat_ru = {'Strength': 'Сила (С)', 'Agility': 'Ловкость (Л)', 'Mastery': 'Мастерство (М)', 'Wit': 'Смекалка (С)'}
                    avg_melted['Показатель'] = avg_melted['Показатель'].map(stat_ru)
                    fig_avg = px.bar(avg_melted, x='Показатель', y='Среднее значение', color='Alignment', barmode='group', color_discrete_map=ALIGN_COLORS)
                    fig_avg.update_layout(margin=dict(t=10, b=10, l=10, r=10), paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font_color='white')
                    st.plotly_chart(fig_avg, use_container_width=True)
                else:
                    st.info("Нет карточек категории «Персонаж» для расчета средних характеристик.")

            st.write("---")
            st.markdown("**🏆 Топ-10 карточек по характеристикам**")
            rank_metric = st.selectbox("Выберите показатель для рейтинга:", [
                "Суммарная мощь (С+Л+М+С)", "Сила (С)", "Ловкость (Л)", "Мастерство (М)", "Смекалка (С)"
            ])
            metric_col_map = {
                "Суммарная мощь (С+Л+М+С)": "Total_Power",
                "Сила (С)": "Strength",
                "Ловкость (Л)": "Agility",
                "Мастерство (М)": "Mastery",
                "Смекалка (С)": "Wit"
            }
            target_m = metric_col_map[rank_metric]
            top_c = df_cards.sort_values(by=target_m, ascending=False).head(10)
            fig_top = px.bar(top_c, x=target_m, y='Name', color='Alignment', orientation='h', color_discrete_map=ALIGN_COLORS, text=target_m)
            fig_top.update_layout(yaxis={'categoryorder':'total ascending'}, margin=dict(t=10, b=10, l=10, r=10), paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font_color='white')
            st.plotly_chart(fig_top, use_container_width=True)


# ==============================================================================
# МОДУЛЬ 3: ФРАКЦИИ, КЛАНЫ И ОРГАНИЗАЦИИ (СИНХРОНИЗАЦИЯ СО СТОЛБЦАМИ ГРАФА)
# ==============================================================================
elif app_mode == "🏛 Фракции и Организации":
    st.title("🏛 Фракции, Кланы и Организации БИГ СИТИ")
    df_fac = st.session_state['factions'].copy()
    if 'Org_Category' not in df_fac.columns:
        df_fac['Org_Category'] = "Группировка"

    all_chars = st.session_state['nodes']['Имя / Название персонажа'].dropna().unique().tolist()
    graph_custom_cols = [c for c in st.session_state['nodes'].columns if c not in ["Имя / Название персонажа", "Заметки"]]

    with st.expander("➕ Создать или отредактировать фракцию / клан / организацию", expanded=df_fac.empty):
        fac_edit = st.selectbox("Выберите организацию:", ["Создать новую запись"] + df_fac['Name'].tolist())
        if fac_edit == "Создать новую запись":
            f_name, f_cat, f_motto, f_leader, f_ideo, f_base, f_inf, f_col, f_desc, f_logo, f_allies, f_enemies = (
                "", "Группировка", "", "Не указано", "", "", 50, "#00d4ff", "", "", "", ""
            )
        else:
            fr = df_fac[df_fac['Name'] == fac_edit].iloc[0]
            f_name = str(fr['Name'])
            f_cat = str(fr.get('Org_Category', 'Группировка')) or "Группировка"
            f_motto, f_leader, f_ideo, f_base = str(fr['Motto']), str(fr['Leader']), str(fr['Ideology']), str(fr['Base_Location'])
            f_inf = int(fr['Influence'])
            f_col = str(fr['Color']) if str(fr['Color']).startswith('#') else "#00d4ff"
            f_desc, f_logo, f_allies, f_enemies = str(fr['Description']), str(fr['Logo_Data']), str(fr['Allies']), str(fr['Enemies'])

        fc1, fc2 = st.columns(2)
        with fc1:
            inp_fname = st.text_input("Название (например: Клан Красного Лотоса):", value=f_name)
            
            base_cat_opts = sorted(list(set(["Группировка", "Фракция", "Клан", "Организация"] + graph_custom_cols)))
            if f_cat not in base_cat_opts:
                base_cat_opts.append(f_cat)
            sel_cat = st.selectbox(
                "🗂 Категория (В какой столбец таблицы Графа добавить):",
                base_cat_opts + ["➕ Создать новый столбец в таблице Графа..."],
                index=base_cat_opts.index(f_cat)
            )
            if sel_cat == "➕ Создать новый столбец в таблице Графа...":
                inp_fcat = st.text_input("Название нового столбца (напр. Клан, Орден, Корпорация):", value="Клан").strip()
            else:
                inp_fcat = sel_cat

            inp_fmotto = st.text_input("Девиз / Слоган:", value=f_motto)
            leader_opts = ["Не указано"] + all_chars
            l_idx = leader_opts.index(f_leader) if f_leader in leader_opts else 0
            inp_fleader = st.selectbox("Лидер / Глава:", leader_opts, index=l_idx)
            inp_fideo = st.text_input("Идеология / Направленность:", value=f_ideo)
            inp_fbase = st.text_input("Штаб-квартира / Главная база:", value=f_base)

        with fc2:
            inp_finf = st.slider("📊 Индекс влияния в городе (0 - 100):", 0, 100, f_inf)
            inp_fcol = st.color_picker("🎨 Фирменный цвет (также раскрасит узлы на Графе!):", f_col)
            inp_fallies = st.text_input("🤝 Союзники (через запятую):", value=f_allies)
            inp_fenemies = st.text_input("⚔️ Враги (через запятую):", value=f_enemies)
            up_logo = st.file_uploader("🛡 Загрузить герб / логотип:", type=["png", "jpg", "jpeg", "webp"], key="fac_logo_up")

        inp_fdesc = st.text_area("История, структура и цели организации (имена героев подсвечиваются):", value=f_desc, height=110)

        fb1, fb2 = st.columns([3, 1])
        with fb1:
            if st.button("💾 Сохранить и синхронизировать с таблицей Графа", type="primary", use_container_width=True):
                if inp_fname.strip() and inp_fcat.strip():
                    clean_fname = inp_fname.strip()
                    clean_fcat = inp_fcat.strip()
                    final_logo = compress_uploaded_image(up_logo, max_size=(300, 300)) if up_logo is not None else f_logo
                    
                    new_f_row = {
                        "Name": clean_fname, "Org_Category": clean_fcat, "Motto": inp_fmotto.strip(),
                        "Leader": inp_fleader, "Ideology": inp_fideo.strip(), "Base_Location": inp_fbase.strip(),
                        "Influence": int(inp_finf), "Color": inp_fcol, "Description": inp_fdesc.strip(),
                        "Logo_Data": final_logo, "Allies": inp_fallies.strip(), "Enemies": inp_fenemies.strip()
                    }
                    
                    for col_req in FAC_COLS:
                        if col_req not in st.session_state['factions'].columns:
                            st.session_state['factions'][col_req] = ""
                    base_f = st.session_state['factions'][FAC_COLS]
                    
                    if fac_edit == "Создать новую запись":
                        st.session_state['factions'] = pd.concat([base_f, pd.DataFrame([new_f_row])], ignore_index=True)
                    else:
                        idx_f = base_f[base_f['Name'] == fac_edit].index[0]
                        for k, v in new_f_row.items():
                            st.session_state['factions'].at[idx_f, k] = v
                        if fac_edit != clean_fname:
                            if clean_fcat in st.session_state['nodes'].columns:
                                st.session_state['nodes'][clean_fcat] = st.session_state['nodes'][clean_fcat].replace(fac_edit, clean_fname)
                                save_sheet("Nodes", st.session_state['nodes'])
                            st.session_state['cards']['Faction'] = st.session_state['cards']['Faction'].replace(fac_edit, clean_fname)
                            save_sheet("Cards", st.session_state['cards'][CARD_COLS])

                    save_sheet("Factions", st.session_state['factions'][FAC_COLS])

                    if clean_fcat not in st.session_state['nodes'].columns:
                        st.session_state['nodes'][clean_fcat] = "Не указано"
                        save_sheet("Nodes", st.session_state['nodes'])

                    opts_list = st.session_state['config']['options'].setdefault(clean_fcat, [])
                    if clean_fname not in opts_list:
                        opts_list.append(clean_fname)
                    st.session_state['config']['colors'].setdefault(clean_fcat, {})[clean_fname] = inp_fcol
                    save_config_to_cloud(st.session_state['config'])

                    st.success(f"«{clean_fname}» сохранена и добавлена в выпадающий список столбца «{clean_fcat}» в таблице персонажей!")
                    st.rerun()
                else:
                    st.warning("Введите название организации и категорию столбца!")
        with fb2:
            if fac_edit != "Создать новую запись":
                if st.button("🗑 Удалить организацию", use_container_width=True):
                    st.session_state['factions'] = df_fac[df_fac['Name'] != fac_edit].reset_index(drop=True)
                    save_sheet("Factions", st.session_state['factions'][FAC_COLS])
                    st.rerun()

    if not df_fac.empty:
        st.subheader("📊 Распределение влияния в БИГ СИТИ")
        fig_inf = px.bar(
            df_fac.sort_values(by='Influence', ascending=True),
            x='Influence', y='Name', orientation='h', text='Influence',
            color='Name', color_discrete_sequence=df_fac.sort_values(by='Influence', ascending=True)['Color'].tolist()
        )
        fig_inf.update_layout(showlegend=False, height=260, margin=dict(t=10, b=10, l=10, r=10), paper_bgcolor='rgba(0,0,0,0)', plot_bgcolor='rgba(0,0,0,0)', font_color='white')
        st.plotly_chart(fig_inf, use_container_width=True)

        st.write("---")
        f_cols = st.columns(2)
        for idx_f, (_, f_row) in enumerate(df_fac.iterrows()):
            f_title = str(f_row['Name']).strip()
            f_cat_col = str(f_row.get('Org_Category', 'Группировка')).strip() or "Группировка"

            nodes_in_fac = []
            check_cols = list(set([f_cat_col, 'Группировка'] + [c for c in st.session_state['nodes'].columns if c not in ['Имя / Название персонажа', 'Заметки']]))
            for c_col in check_cols:
                if c_col in st.session_state['nodes'].columns:
                    matched_chars = st.session_state['nodes'][
                        st.session_state['nodes'][c_col].astype(str).str.strip().str.lower() == f_title.lower()
                    ]['Имя / Название персонажа'].tolist()
                    nodes_in_fac.extend(matched_chars)

            cards_in_fac = st.session_state['cards'][
                st.session_state['cards']['Faction'].astype(str).str.strip().str.lower() == f_title.lower()
            ]['Name'].tolist()
            
            leader_name = str(f_row.get('Leader', '')).strip()
            if leader_name and leader_name != "Не указано":
                nodes_in_fac.append(leader_name)

            all_members = sorted(list(set(nodes_in_fac + cards_in_fac)))

            locs_in_fac = st.session_state['locations'][
                st.session_state['locations']['Controlling_Faction'].astype(str).str.strip().str.lower() == f_title.lower()
            ]['Name'].tolist()

            with f_cols[idx_f % 2]:
                st.markdown(f"""
                <div class="game-card" style="border-top: 4px solid {html.escape(str(f_row['Color']))};">
                    <div class="game-card-header">
                        <span style="color:{html.escape(str(f_row['Color']))}; font-weight:bold;">ВЛИЯНИЕ: {int(f_row['Influence'])} / 100</span>
                        <span class="game-card-type">Столбец: {html.escape(f_cat_col)} • {html.escape(str(f_row['Ideology'] or 'Организация'))}</span>
                    </div>
                    <div class="game-card-title">{html.escape(f_title)}</div>
                    <div style="font-style:italic; color:#ffcc00; font-size:13px; margin-bottom:8px;">«{html.escape(str(f_row['Motto']))}»</div>
                </div>
                """, unsafe_allow_html=True)
                if f_row['Logo_Data']:
                    st.image(f_row['Logo_Data'], width=140)
                st.progress(min(max(int(f_row['Influence']), 0), 100) / 100.0)
                st.markdown(f"👑 **Лидер:** {f_row['Leader']} | 📍 **База:** {f_row['Base_Location'] or 'Не указана'}")
                if f_row['Allies']: st.markdown(f"🤝 **Союзники:** {f_row['Allies']}")
                if f_row['Enemies']: st.markdown(f"⚔️ **Враги:** {f_row['Enemies']}")
                st.markdown(f"👥 **Участники ({len(all_members)}):** {', '.join(all_members) if all_members else 'Пока нет привязанных героев'}")
                st.markdown(f"🗺 **Подконтрольные территории ({len(locs_in_fac)}):** {', '.join(locs_in_fac) if locs_in_fac else 'Нет точек на карте'}")
                if f_row['Description']:
                    hl_fdesc, _ = render_text_with_spoilers_and_characters(
                        str(f_row['Description']), CHAR_MATCHERS, CHAR_POPUP_SPANS, enable_spoilers=False
                    )
                    st.markdown(f"<div class='script-body' style='font-size:14px;'>{hl_fdesc}</div>", unsafe_allow_html=True)
                st.write("---")


# ==============================================================================
# МОДУЛЬ 4: КАРТА МЕСТНОСТИ (КЛИК МЫШЬЮ + УДАЛЕНИЕ ПЕРСОНАЖЕЙ И ЛОКАЦИЙ)
# ==============================================================================
elif app_mode == "🗺 Карта Местности":
    st.title("🗺 Интерактивная Карта: Размещение и Управление")
    df_loc = st.session_state['locations'].copy()
    all_chars = st.session_state['nodes']['Имя / Название персонажа'].dropna().unique().tolist()
    all_scripts = st.session_state['scripts']['Title'].dropna().unique().tolist()
    all_factions = ["Нейтральная территория"] + st.session_state['factions']['Name'].dropna().unique().tolist()

    if 'map_clicked_x' not in st.session_state:
        st.session_state['map_clicked_x'] = 50.0
    if 'map_clicked_y' not in st.session_state:
        st.session_state['map_clicked_y'] = 50.0
    if 'map_selected_item' not in st.session_state:
        st.session_state['map_selected_item'] = "➕ Создать в точке прицела 🎯"

    existing_regions = sorted(list(set(["Глобальная карта (БИГ СИТИ)"] + [r for r in df_loc['Region'].unique().tolist() if str(r).strip()])))
    col_reg1, col_reg2 = st.columns([2, 1.2])
    with col_reg1:
        active_region = st.selectbox("🌍 Выберите уровень карты / Регион:", existing_regions)
        new_reg_input = st.text_input("➕ Или введите название нового региона / плана здания:", placeholder="Например: Район Доков / Этаж 1")
        if new_reg_input.strip():
            active_region = new_reg_input.strip()

    with col_reg2:
        with st.expander("🖼 Фоновая картинка этой карты", expanded=not bool(st.session_state['config'].get('region_maps', {}).get(active_region))):
            up_reg_bg = st.file_uploader("Загрузить изображение (PNG/JPG):", type=["png", "jpg", "jpeg", "webp"], key="reg_bg_up")
            bg_url_inp = st.text_input("ИЛИ прямая ссылка на картинку карты:", key="reg_bg_url")
            if st.button("💾 Установить фон карты", use_container_width=True):
                if up_reg_bg is not None:
                    b64_map = compress_uploaded_image(up_reg_bg, max_size=(900, 650), quality=70)
                    st.session_state['config'].setdefault('region_maps', {})[active_region] = b64_map
                    save_config_to_cloud(st.session_state['config'])
                    st.success("Картинка карты сохранена!")
                    st.rerun()
                elif bg_url_inp.strip():
                    st.session_state['config'].setdefault('region_maps', {})[active_region] = bg_url_inp.strip()
                    save_config_to_cloud(st.session_state['config'])
                    st.success("Ссылка на карту сохранена!")
                    st.rerun()

            if st.session_state['config'].get('region_maps', {}).get(active_region):
                if st.button("🗑 Удалить фон региона", use_container_width=True):
                    del st.session_state['config']['region_maps'][active_region]
                    save_config_to_cloud(st.session_state['config'])
                    st.rerun()

    st.info("👆 **Управление:** Кликните в пустое место карты, чтобы поставить **прицел 🎯**, ИЛИ кликните прямо по **существующему маркеру** на карте, чтобы изменить или удалить его!")

    map_col, form_col = st.columns([1.85, 1.15])
    reg_df = df_loc[df_loc['Region'] == active_region].copy()

    with map_col:
        fig_map = go.Figure()

        bg_map_data = st.session_state['config'].get('region_maps', {}).get(active_region)
        if bg_map_data:
            fig_map.add_layout_image(
                dict(
                    source=bg_map_data,
                    xref="x", yref="y",
                    x=0, y=100,
                    sizex=100, sizey=100,
                    sizing="stretch",
                    opacity=0.92,
                    layer="below"
                )
            )

        grid_x, grid_y = [], []
        for gx in range(0, 102, 2):
            for gy in range(0, 102, 2):
                grid_x.append(gx)
                grid_y.append(gy)

        fig_map.add_trace(go.Scatter(
            x=grid_x,
            y=grid_y,
            mode='markers',
            name='Сетка клика',
            showlegend=False,
            marker=dict(size=14, color='rgba(0, 0, 0, 0.005)'),
            hoverinfo='x+y',
            hovertemplate="🎯 Кликните для выбора координат: X=%{x}, Y=%{y}<extra></extra>"
        ))

        if not reg_df.empty:
            type_style_map = {
                "👤 Персонаж на карте": ("circle", "#00d4ff", 18),
                "Городской объект": ("octagon", "#ffcc00", 16),
                "Штаб / База": ("star", "#ff4d4d", 19),
                "Район / Зона": ("hexagon", "#33cc33", 17),
                "Подземелье / Секрет": ("diamond", "#9933ff", 16),
                "Порт / Узел": ("square", "#ff9933", 15),
                "Руины / Аномалия": ("cross", "#ff66b2", 16)
            }

            for l_type_val, grp in reg_df.groupby('Loc_Type'):
                sym, col_m, sz = type_style_map.get(str(l_type_val), ("octagon", "#00d4ff", 16))
                fig_map.add_trace(go.Scatter(
                    x=grp['X'], y=grp['Y'],
                    mode='markers+text',
                    name=str(l_type_val),
                    text=grp['Name'],
                    textposition='top center',
                    textfont=dict(size=12, color='#ffffff', family='Courier New'),
                    marker=dict(size=sz, symbol=sym, color=col_m, line=dict(width=2, color='#121212')),
                    customdata=grp[['Controlling_Faction', 'Linked_Characters', 'Linked_Scripts', 'Description', 'Name']],
                    hovertemplate=(
                        "<b>%{text}</b> (" + str(l_type_val) + ")<br>"
                        "Фракция/Контроль: <b>%{customdata[0]}</b><br>"
                        "Координаты: X=%{x}, Y=%{y}<br>"
                        "Герои: %{customdata[1]}<br>"
                        "<i>%{customdata[3]}</i><extra>Кликните для выбора/удаления</extra>"
                    )
                ))

        fig_map.add_trace(go.Scatter(
            x=[st.session_state['map_clicked_x']],
            y=[st.session_state['map_clicked_y']],
            mode='markers+text',
            name='🎯 Выбранная точка',
            text=[f"🎯 ({st.session_state['map_clicked_x']}, {st.session_state['map_clicked_y']})"],
            textposition='bottom center',
            textfont=dict(size=11, color='#ffd700', family='Courier New'),
            marker=dict(size=20, symbol='cross-thin', color='#ffd700', line=dict(width=3, color='#ffd700')),
            hoverinfo='skip'
        ))

        fig_map.update_layout(
            title=dict(text=f"🗺 {active_region}", font=dict(size=16, color="#00d4ff")),
            height=620,
            margin=dict(l=10, r=10, t=40, b=10),
            xaxis=dict(range=[0, 100], showgrid=True, gridcolor='rgba(255,255,255,0.07)', zeroline=False, fixedrange=False, title="Ось X"),
            yaxis=dict(range=[0, 100], showgrid=True, gridcolor='rgba(255,255,255,0.07)', zeroline=False, fixedrange=False, title="Ось Y"),
            paper_bgcolor='#14141f',
            plot_bgcolor='#181828',
            font=dict(color='white'),
            legend=dict(orientation="h", yanchor="bottom", y=1.01, xanchor="right", x=1),
            clickmode='event+select',
            dragmode='pan'
        )

        map_event = st.plotly_chart(
            fig_map,
            use_container_width=True,
            on_select="rerun",
            selection_mode="points",
            key=f"interactive_map_{active_region}"
        )

        if map_event and "selection" in map_event and map_event["selection"].get("points"):
            pt = map_event["selection"]["points"][0]
            new_x = float(pt.get("x", st.session_state['map_clicked_x']))
            new_y = float(pt.get("y", st.session_state['map_clicked_y']))

            clicked_item_name = None
            if not reg_df.empty:
                matched = reg_df[(reg_df['X'] == new_x) & (reg_df['Y'] == new_y)]
                if not matched.empty:
                    clicked_item_name = str(matched.iloc[0]['Name'])

            coords_changed = (new_x, new_y) != (st.session_state['map_clicked_x'], st.session_state['map_clicked_y'])
            item_changed = clicked_item_name is not None and clicked_item_name != st.session_state.get('map_selected_item')

            if coords_changed or item_changed:
                st.session_state['map_clicked_x'] = new_x
                st.session_state['map_clicked_y'] = new_y
                if clicked_item_name:
                    st.session_state['map_selected_item'] = clicked_item_name
                st.rerun()

    with form_col:
        st.subheader("🎯 Панель управления точкой")
        st.markdown(f"**Координаты прицела:** `X = {st.session_state['map_clicked_x']}` | `Y = {st.session_state['map_clicked_y']}`")

        place_mode = st.radio(
            "Режим работы:",
            ["📍 Локации и Маркеры (Все)", "👤 Быстро поставить Персонажа"],
            horizontal=True
        )

        if place_mode == "👤 Быстро поставить Персонажа":
            if not all_chars:
                st.warning("В базе пока нет персонажей.")
            else:
                sel_char_map = st.selectbox("Выберите персонажа:", all_chars, key="map_char_sel")
                marker_name = f"👤 {sel_char_map}"
                char_info = CHAR_LOOKUP.get(sel_char_map, {})
                default_fac = char_info.get("faction", "Нейтральная территория") if char_info.get("faction") else "Нейтральная территория"

                existing_char_mask = (df_loc['Name'] == marker_name) & (df_loc['Region'] == active_region)
                is_already_on_map = existing_char_mask.any()

                if is_already_on_map:
                    ex_row = df_loc[existing_char_mask].iloc[0]
                    st.caption(f"✅ Сейчас на этой карте в точке `X={ex_row['X']}, Y={ex_row['Y']}`")
                    default_note = str(ex_row['Description'])
                else:
                    st.caption("⚪ Этого персонажа пока нет на данной карте.")
                    default_note = "Находится здесь"

                char_note = st.text_input("Статус / Что делает в этой точке:", value=default_note)

                if st.button("📌 Поставить / Переместить в точку 🎯", type="primary", use_container_width=True):
                    new_c_loc = {
                        "Name": marker_name,
                        "Region": active_region,
                        "Loc_Type": "👤 Персонаж на карте",
                        "Controlling_Faction": default_fac,
                        "X": float(st.session_state['map_clicked_x']),
                        "Y": float(st.session_state['map_clicked_y']),
                        "Description": char_note.strip(),
                        "Map_Image": "",
                        "Linked_Characters": sel_char_map,
                        "Linked_Scripts": ""
                    }
                    base_l = st.session_state['locations'][LOC_COLS]
                    if is_already_on_map:
                        idx_m = base_l[existing_char_mask].index[0]
                        for k, v in new_c_loc.items():
                            st.session_state['locations'].at[idx_m, k] = v
                    else:
                        st.session_state['locations'] = pd.concat([base_l, pd.DataFrame([new_c_loc])], ignore_index=True)

                    save_sheet("Locations", st.session_state['locations'][LOC_COLS])
                    st.success(f"{sel_char_map} размещен в ({st.session_state['map_clicked_x']}, {st.session_state['map_clicked_y']})!")
                    st.rerun()

                if is_already_on_map:
                    if st.button(f"🗑 Убрать «{sel_char_map}» с этой карты", use_container_width=True):
                        st.session_state['locations'] = df_loc[~existing_char_mask].reset_index(drop=True)
                        save_sheet("Locations", st.session_state['locations'][LOC_COLS])
                        st.success(f"{sel_char_map} убран с карты!")
                        st.rerun()

        else:
            item_options = ["➕ Создать в точке прицела 🎯"] + reg_df['Name'].tolist()
            default_idx = item_options.index(st.session_state['map_selected_item']) if st.session_state.get('map_selected_item') in item_options else 0
            loc_edit = st.selectbox("Выберите элемент на карте (или кликните по нему мышкой):", item_options, index=default_idx)
            st.session_state['map_selected_item'] = loc_edit

            if loc_edit == "➕ Создать в точке прицела 🎯":
                l_name, l_type, l_fac, l_desc, l_chars, l_scrs = (
                    "", "Городской объект", "Нейтральная территория", "", [], []
                )
                cur_x, cur_y = st.session_state['map_clicked_x'], st.session_state['map_clicked_y']
            else:
                lr = reg_df[reg_df['Name'] == loc_edit].iloc[0]
                l_name, l_type, l_fac = str(lr['Name']), str(lr['Loc_Type']), str(lr['Controlling_Faction'])
                l_desc = str(lr['Description'])
                l_chars = [x.strip() for x in str(lr['Linked_Characters']).split(',') if x.strip() in all_chars]
                l_scrs = [x.strip() for x in str(lr['Linked_Scripts']).split(',') if x.strip() in all_scripts]
                cur_x, cur_y = float(lr['X']), float(lr['Y'])

            inp_lname = st.text_input("Название объекта / локации / маркера:", value=l_name)
            type_l_opts = ["Городской объект", "Штаб / База", "Район / Зона", "Подземелье / Секрет", "Порт / Узел", "Руины / Аномалия", "👤 Персонаж на карте"]
            inp_ltype = st.selectbox("Тип точки:", type_l_opts, index=type_l_opts.index(l_type) if l_type in type_l_opts else 0)
            fac_opts = all_factions if l_fac in all_factions else all_factions + [l_fac]
            inp_lfac = st.selectbox("Контролирующая фракция:", fac_opts, index=fac_opts.index(l_fac))

            move_to_crosshair = st.checkbox(
                f"🎯 Переместить в текущие координаты прицела ({st.session_state['map_clicked_x']}, {st.session_state['map_clicked_y']})",
                value=(loc_edit == "➕ Создать в точке прицела 🎯")
            )
            final_x = st.session_state['map_clicked_x'] if move_to_crosshair else cur_x
            final_y = st.session_state['map_clicked_y'] if move_to_crosshair else cur_y

            inp_lchars = st.multiselect("👥 Привязанные персонажи:", all_chars, default=l_chars)
            inp_lscrs = st.multiselect("🎬 Связанные сценарии:", all_scripts, default=l_scrs)
            inp_ldesc = st.text_area("Описание / Лор объекта:", value=l_desc, height=75)

            b_save, b_del = st.columns([1.6, 1.4])
            with b_save:
                if st.button("💾 Сохранить", type="primary", use_container_width=True):
                    if inp_lname.strip():
                        new_l_row = {
                            "Name": inp_lname.strip(),
                            "Region": active_region,
                            "Loc_Type": inp_ltype,
                            "Controlling_Faction": inp_lfac,
                            "X": float(final_x),
                            "Y": float(final_y),
                            "Description": inp_ldesc.strip(),
                            "Map_Image": "",
                            "Linked_Characters": ", ".join(inp_lchars),
                            "Linked_Scripts": ", ".join(inp_lscrs)
                        }
                        base_l = st.session_state['locations'][LOC_COLS]
                        if loc_edit == "➕ Создать в точке прицела 🎯":
                            st.session_state['locations'] = pd.concat([base_l, pd.DataFrame([new_l_row])], ignore_index=True)
                        else:
                            idx_l = base_l[(base_l['Name'] == loc_edit) & (base_l['Region'] == active_region)].index[0]
                            for k, v in new_l_row.items():
                                st.session_state['locations'].at[idx_l, k] = v
                        save_sheet("Locations", st.session_state['locations'][LOC_COLS])
                        st.session_state['map_selected_item'] = inp_lname.strip()
                        st.success("Точка сохранена!")
                        st.rerun()
                    else:
                        st.warning("Введите название!")
            with b_del:
                if loc_edit != "➕ Создать в точке прицела 🎯":
                    if st.button("🗑 Удалить с карты", use_container_width=True):
                        st.session_state['locations'] = df_loc[~((df_loc['Name'] == loc_edit) & (df_loc['Region'] == active_region))].reset_index(drop=True)
                        save_sheet("Locations", st.session_state['locations'][LOC_COLS])
                        st.session_state['map_selected_item'] = "➕ Создать в точке прицела 🎯"
                        st.success(f"«{loc_edit}» удален с карты!")
                        st.rerun()

    st.write("---")
    tab_loc_cards, tab_loc_table = st.tabs([
        f"📍 Все объекты и персонажи региона «{active_region}» ({len(reg_df)})",
        "📋 Табличный редактор всех локаций (Массовое удаление)"
    ])

    with tab_loc_cards:
        if reg_df.empty:
            st.caption("На этой карте пока не размещено ни одного объекта или персонажа.")
        else:
            l_cols = st.columns(2)
            for i_l, (orig_idx, l_item) in enumerate(reg_df.iterrows()):
                with l_cols[i_l % 2]:
                    st.markdown(f"""
                    <div class="game-card">
                        <div class="game-card-header">
                            <span style="color:#ffcc00; font-size:12px;">🏛 {html.escape(str(l_item['Controlling_Faction']))}</span>
                            <span class="game-card-type">{html.escape(str(l_item['Loc_Type']))} (X:{l_item['X']}, Y:{l_item['Y']})</span>
                        </div>
                        <div class="game-card-title">{html.escape(str(l_item['Name']))}</div>
                    </div>
                    """, unsafe_allow_html=True)
                    if l_item['Linked_Characters']:
                        st.caption(f"👥 **Персонажи:** {l_item['Linked_Characters']}")
                    if l_item['Linked_Scripts']:
                        st.caption(f"🎬 **Сценарии:** {l_item['Linked_Scripts']}")
                    if l_item['Description']:
                        hl_ldesc, _ = render_text_with_spoilers_and_characters(
                            str(l_item['Description']), CHAR_MATCHERS, CHAR_POPUP_SPANS, enable_spoilers=False
                        )
                        st.markdown(f"<div class='script-body' style='font-size:14px;'>{hl_ldesc}</div>", unsafe_allow_html=True)

                    act_c1, act_c2 = st.columns(2)
                    with act_c1:
                        if st.button("✏️ Выбрать на карте", key=f"sel_loc_{orig_idx}", use_container_width=True):
                            st.session_state['map_selected_item'] = str(l_item['Name'])
                            st.session_state['map_clicked_x'] = float(l_item['X'])
                            st.session_state['map_clicked_y'] = float(l_item['Y'])
                            st.rerun()
                    with act_c2:
                        if st.button("🗑 Убрать с карты", key=f"del_loc_{orig_idx}", use_container_width=True):
                            st.session_state['locations'] = df_loc.drop(index=orig_idx).reset_index(drop=True)
                            save_sheet("Locations", st.session_state['locations'][LOC_COLS])
                            if st.session_state.get('map_selected_item') == str(l_item['Name']):
                                st.session_state['map_selected_item'] = "➕ Создать в точке прицела 🎯"
                            st.rerun()

    with tab_loc_table:
        st.caption("💡 Здесь можно выделить галочками слева сразу несколько строк (локаций или персонажей) и нажать клавишу `Delete` или иконку корзины в правом верхнем углу таблицы:")
        edited_locs = st.data_editor(
            st.session_state['locations'][LOC_COLS],
            num_rows="dynamic",
            use_container_width=True,
            height=400,
            key="loc_table_editor"
        )
        if st.button("💾 Сохранить изменения в таблице локаций", type="primary"):
            clean_locs = edited_locs.dropna(subset=['Name']).reset_index(drop=True)
            st.session_state['locations'] = clean_locs
            save_sheet("Locations", clean_locs[LOC_COLS])
            st.success("Таблица локаций синхронизирована с облаком!")
            st.rerun()


# ==============================================================================
# МОДУЛЬ 5: ХРОНОЛОГИЯ ВСЕЛЕННОЙ (TIMELINE ПО ЭПОХАМ И ИНДЕКСАМ)
# ==============================================================================
elif app_mode == "⏳ Хронология Вселенной":
    st.title("⏳ Хронология и Летопись Вселенной БИГ СИТИ")
    df_time = st.session_state['timeline'].copy()
    all_chars = st.session_state['nodes']['Имя / Название персонажа'].dropna().unique().tolist()
    all_locs = ["Не указано"] + st.session_state['locations']['Name'].dropna().unique().tolist()
    all_scrs = ["Нет привязки"] + st.session_state['scripts']['Title'].dropna().unique().tolist()

    with st.expander("➕ Добавить или изменить историческое событие", expanded=df_time.empty):
        ev_options = ["Добавить новое событие"] + [f"[{r['Order_Index']}] {r['Title']}" for _, r in df_time.sort_values(by='Order_Index').iterrows()]
        ev_edit = st.selectbox("Выберите событие:", ev_options)
        if ev_edit == "Добавить новое событие":
            next_idx = int(df_time['Order_Index'].max()) + 10 if not df_time.empty else 10
            t_idx, t_date, t_era, t_title, t_desc, t_parts, t_loc, t_scr = (
                next_idx, "", "Эпоха Основания", "", "", [], "Не указано", "Нет привязки"
            )
        else:
            chosen_title = ev_edit.split("] ", 1)[1]
            tr = df_time[df_time['Title'] == chosen_title].iloc[0]
            t_idx = int(tr['Order_Index'])
            t_date, t_era, t_title, t_desc = str(tr['Date_Label']), str(tr['Era']), str(tr['Title']), str(tr['Description'])
            t_parts = [x.strip() for x in str(tr['Participants']).split(',') if x.strip() in all_chars]
            t_loc = str(tr['Location']) if str(tr['Location']) in all_locs else "Не указано"
            t_scr = str(tr['Linked_Script']) if str(tr['Linked_Script']) in all_scrs else "Нет привязки"

        tc1, tc2 = st.columns(2)
        with tc1:
            inp_tidx = st.number_input("🔢 Порядковый индекс на шкале времени (для сортировки):", value=t_idx, step=5)
            inp_tdate = st.text_input("📅 Дата / Временная метка (напр. «За 15 лет до Катаклизма» или «2084 год»):", value=t_date)
            inp_tera = st.text_input("🏛 Эпоха / Исторический период:", value=t_era)
            inp_ttitle = st.text_input("📌 Название события:", value=t_title)
        with tc2:
            inp_tparts = st.multiselect("👥 Ключевые участники:", all_chars, default=t_parts)
            inp_tloc = st.selectbox("📍 Локация события:", all_locs, index=all_locs.index(t_loc))
            inp_tscr = st.selectbox("🎬 Связанный сценарий:", all_scrs, index=all_scrs.index(t_scr))

        inp_tdesc = st.text_area("Подробное описание события (с авто-подсветкой имён героев):", value=t_desc, height=110)

        tb1, tb2 = st.columns([3, 1])
        with tb1:
            if st.button("💾 Сохранить событие в летопись", type="primary", use_container_width=True):
                if inp_ttitle.strip():
                    new_t_row = {
                        "Order_Index": int(inp_tidx),
                        "Date_Label": inp_tdate.strip() or f"Индекс {inp_tidx}",
                        "Era": inp_tera.strip() or "Общая эпоха",
                        "Title": inp_ttitle.strip(),
                        "Description": inp_tdesc.strip(),
                        "Participants": ", ".join(inp_tparts),
                        "Location": inp_tloc if inp_tloc != "Не указано" else "",
                        "Linked_Script": inp_tscr if inp_tscr != "Нет привязки" else ""
                    }
                    base_t = st.session_state['timeline'][TIME_COLS]
                    if ev_edit == "Добавить новое событие":
                        st.session_state['timeline'] = pd.concat([base_t, pd.DataFrame([new_t_row])], ignore_index=True)
                    else:
                        chosen_title = ev_edit.split("] ", 1)[1]
                        idx_t = base_t[base_t['Title'] == chosen_title].index[0]
                        for k, v in new_t_row.items():
                            st.session_state['timeline'].at[idx_t, k] = v
                    save_sheet("Timeline", st.session_state['timeline'][TIME_COLS])
                    st.success("Событие сохранено в Хронологию!")
                    st.rerun()
                else:
                    st.warning("Введите название события!")
        with tb2:
            if ev_edit != "Добавить новое событие":
                if st.button("🗑 Удалить событие", use_container_width=True):
                    chosen_title = ev_edit.split("] ", 1)[1]
                    st.session_state['timeline'] = df_time[df_time['Title'] != chosen_title].reset_index(drop=True)
                    save_sheet("Timeline", st.session_state['timeline'][TIME_COLS])
                    st.rerun()

    if not df_time.empty:
        eras = ["Все эпохи"] + sorted([e for e in df_time['Era'].unique().tolist() if str(e).strip()])
        sel_era = st.selectbox("Фильтр по эпохе:", eras)
        view_time = df_time.sort_values(by='Order_Index')
        if sel_era != "Все эпохи":
            view_time = view_time[view_time['Era'] == sel_era]

        fig_tl = go.Figure()
        fig_tl.add_hline(y=0, line_width=3, line_color="#00d4ff")
        for era_name, grp in view_time.groupby('Era', sort=False):
            fig_tl.add_trace(go.Scatter(
                x=grp['Order_Index'],
                y=[0] * len(grp),
                mode='markers+text',
                name=str(era_name),
                text=grp['Title'],
                textposition="top center",
                marker=dict(size=15, symbol="diamond", line=dict(width=2, color="#ffffff")),
                customdata=grp[['Date_Label', 'Location', 'Participants']],
                hovertemplate="<b>%{customdata[0]} — %{text}</b><br>Эпоха: " + str(era_name) + "<br>Место: %{customdata[1]}<br>Участники: %{customdata[2]}<extra></extra>"
            ))
        fig_tl.update_layout(
            title=dict(text="🕰 Визуальная шкала хронологии (по числовым индексам)", font=dict(size=14, color="#00d4ff")),
            height=220,
            margin=dict(l=20, r=20, t=40, b=25),
            paper_bgcolor='#161622', plot_bgcolor='#161622', font=dict(color='white'),
            xaxis=dict(title="Хронологический индекс (Order_Index)", gridcolor='#26263a'),
            yaxis=dict(visible=False, range=[-0.5, 0.7]),
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
        )
        st.plotly_chart(fig_tl, use_container_width=True)

        st.write("---")
        for _, t_row in view_time.iterrows():
            st.markdown(f"""
            <div class="game-card" style="border-left: 4px solid #00d4ff;">
                <div class="game-card-header">
                    <span style="color:#ffcc00; font-weight:bold;">📅 {html.escape(str(t_row['Date_Label']))} (Индекс: {int(t_row['Order_Index'])})</span>
                    <span class="game-card-type">{html.escape(str(t_row['Era']))}</span>
                </div>
                <div class="game-card-title">{html.escape(str(t_row['Title']))}</div>
            </div>
            """, unsafe_allow_html=True)
            meta_t = []
            if t_row['Location']: meta_t.append(f"📍 **Локация:** {t_row['Location']}")
            if t_row['Participants']: meta_t.append(f"👥 **Участники:** {t_row['Participants']}")
            if t_row['Linked_Script']: meta_t.append(f"🎬 **Сценарий:** {t_row['Linked_Script']}")
            if meta_t:
                st.caption(" | ".join(meta_t))
            if t_row['Description']:
                hl_tdesc, _ = render_text_with_spoilers_and_characters(
                    str(t_row['Description']), CHAR_MATCHERS, CHAR_POPUP_SPANS, enable_spoilers=False
                )
                st.markdown(f"<div class='script-body'>{hl_tdesc}</div>", unsafe_allow_html=True)


# ==============================================================================
# МОДУЛЬ 6: АРХИВЫ И ЛОР (С АВТО-ПОДСВЕТКОЙ ИМЁН И ЗАГРУЗКОЙ ФАЙЛОВ)
# ==============================================================================
elif app_mode == "📚 Архивы и Лор":
    st.title("📚 Архивы БИГ СИТИ")
    df_arch = st.session_state['archives']

    col_left, col_right = st.columns([1, 2.5])

    with col_left:
        st.subheader("🗂 Оглавление")
        if st.button("➕ Создать / Загрузить документ", use_container_width=True, type="primary"):
            st.session_state['selected_doc'] = "__NEW__"

        categories = ["Все"] + sorted([c for c in df_arch['Category'].unique().tolist() if c])
        filter_cat = st.selectbox("Фильтр по разделу:", categories)

        filtered_arch = df_arch if filter_cat == "Все" else df_arch[df_arch['Category'] == filter_cat]
        doc_titles = filtered_arch['Title'].tolist()

        if doc_titles:
            chosen_doc = st.radio("Выберите запись:", doc_titles)
            if st.session_state.get('selected_doc') != "__NEW__":
                st.session_state['selected_doc'] = chosen_doc
        else:
            st.info("В этом разделе пока нет записей.")

    with col_right:
        if st.session_state.get('selected_doc') == "__NEW__" or df_arch.empty:
            st.subheader("✍️ Новый архивный документ")
            
            uploaded_doc = st.file_uploader("📂 Загрузить готовый файл с компьютера (.docx, .txt, .md):", type=["docx", "txt", "md"], key="arch_upload")
            init_title, init_content = "", ""
            if uploaded_doc is not None:
                init_title, init_content = read_uploaded_document(uploaded_doc)
                st.success(f"Файл «{uploaded_doc.name}» успешно прочитан! Проверьте текст ниже и нажмите «Сохранить в Архив».")

            new_title = st.text_input("Название документа:", value=init_title)
            new_cat = st.text_input("Раздел / Том (например: Лор, Правила, Локации):", value="Общий Лор")
            new_content = st.text_area("Текст документа (все имена персонажей подсвечиваются автоматически):", value=init_content, height=400)
            
            c1, c2 = st.columns(2)
            with c1:
                if st.button("💾 Сохранить в Архив", type="primary", use_container_width=True):
                    if new_title.strip():
                        new_row = pd.DataFrame([{"Title": new_title.strip(), "Category": new_cat.strip(), "Content": new_content}])
                        st.session_state['archives'] = pd.concat([df_arch, new_row], ignore_index=True)
                        save_sheet("Archives", st.session_state['archives'])
                        st.session_state['selected_doc'] = new_title.strip()
                        st.success("Документ сохранен!")
                        st.rerun()
                    else:
                        st.warning("Введите название документа!")
            with c2:
                if st.button("Отмена", use_container_width=True):
                    st.session_state['selected_doc'] = doc_titles[0] if doc_titles else None
                    st.rerun()

        elif st.session_state.get('selected_doc') in df_arch['Title'].values:
            curr_title = st.session_state['selected_doc']
            doc_row = df_arch[df_arch['Title'] == curr_title].iloc[0]

            mode_read, mode_edit = st.tabs(["📖 Интерактивное чтение", "✏️ Редактировать / Заменить файлом"])
            with mode_read:
                st.caption(f"Раздел: **{doc_row['Category']}**")
                st.markdown(f"## {doc_row['Title']}")
                
                rendered_arch_html, arch_found_chars = render_text_with_spoilers_and_characters(
                    str(doc_row['Content']), CHAR_MATCHERS, CHAR_POPUP_SPANS, enable_spoilers=True
                )
                render_character_popover_bar(arch_found_chars, CHAR_LOOKUP, key_prefix="arch")
                st.write("---")
                st.markdown(f"<div class='script-body'>{rendered_arch_html}</div>", unsafe_allow_html=True)

            with mode_edit:
                replace_doc = st.file_uploader("📂 Заменить текст из файла (.docx, .txt, .md):", type=["docx", "txt", "md"], key="arch_replace")
                curr_content_val = doc_row['Content']
                if replace_doc is not None:
                    _, curr_content_val = read_uploaded_document(replace_doc)
                    st.info("Текст из файла подгружен в поле ниже. Нажмите «Обновить документ», чтобы сохранить.")

                ed_title = st.text_input("Название:", value=doc_row['Title'])
                ed_cat = st.text_input("Раздел:", value=doc_row['Category'])
                ed_content = st.text_area("Содержание:", value=curr_content_val, height=420)
                
                b1, b2 = st.columns([3, 1])
                with b1:
                    if st.button("💾 Обновить документ", type="primary", use_container_width=True):
                        idx = df_arch[df_arch['Title'] == curr_title].index[0]
                        st.session_state['archives'].at[idx, 'Title'] = ed_title.strip()
                        st.session_state['archives'].at[idx, 'Category'] = ed_cat.strip()
                        st.session_state['archives'].at[idx, 'Content'] = ed_content
                        save_sheet("Archives", st.session_state['archives'])
                        st.session_state['selected_doc'] = ed_title.strip()
                        st.success("Обновлено!")
                        st.rerun()
                with b2:
                    if st.button("🗑 Удалить", use_container_width=True):
                        st.session_state['archives'] = df_arch[df_arch['Title'] != curr_title].reset_index(drop=True)
                        save_sheet("Archives", st.session_state['archives'])
                        st.session_state['selected_doc'] = None
                        st.rerun()


# ==============================================================================
# МОДУЛЬ 7: СЦЕНАРИИ И СЮЖЕТ (СО СПОЙЛЕРАМИ И ИНТЕРАКТИВНЫМИ ИМЕНАМИ ГЕРОЕВ)
# ==============================================================================
elif app_mode == "🎬 Сценарии и Сюжет":
    st.title("🎬 Сценарии и Сюжетные Линии")
    df_scr = st.session_state['scripts']

    col_s_left, col_s_right = st.columns([1, 2.5])

    with col_s_left:
        st.subheader("📑 Сюжетные ветки")
        if st.button("➕ Новый / Загрузить сценарий", use_container_width=True, type="primary"):
            st.session_state['selected_script'] = "__NEW_SCRIPT__"

        scr_titles = df_scr['Title'].tolist()
        if scr_titles:
            chosen_scr = st.radio("Выберите сценарий:", scr_titles)
            if st.session_state.get('selected_script') != "__NEW_SCRIPT__":
                st.session_state['selected_script'] = chosen_scr
        else:
            st.info("Сценариев пока нет.")

        with st.expander("💡 Подсказка по разметке"):
            st.markdown("""
            * `# Заголовок` — **главный скрытый блок** (1 уровень).
            * `## Подзаголовок` — **вложенный подзаголовок** (2 уровень).
            * `### Деталь` — **глубокий спойлер** (3 уровень).
            * **Имена любых персонажей** из базы подсвечиваются в тексте автоматически! Наведите на имя курсор (или нажмите на телефоне), чтобы увидеть его карточку.
            """)

    with col_s_right:
        default_template = (
            "Вступление к сценарию (видно сразу).\n\n"
            "# Акт 1. Тень над Городом\n"
            "Общее описание первого акта.\n\n"
            "## Сцена 1. Встреча в порту\n"
            "Диалог героев у старого склада.\n\n"
            "### Секретная улика (Спойлер)\n"
            "Если игроки обыщут ящик, они найдут медальон."
        )

        if st.session_state.get('selected_script') == "__NEW_SCRIPT__" or df_scr.empty:
            st.subheader("✍️ Создание или загрузка сценария")
            
            up_scr = st.file_uploader("📂 Загрузить готовый сценарий с компьютера (.docx, .txt, .md):", type=["docx", "txt", "md"], key="scr_upload")
            s_init_title, s_init_content = "", default_template
            if up_scr is not None:
                s_init_title, s_init_content = read_uploaded_document(up_scr)
                st.success(f"Сценарий «{up_scr.name}» загружен! Проверьте разметку и нажмите «Сохранить сценарий».")

            s_title = st.text_input("Название сценария / Главы:", value=s_init_title)
            c_a1, c_a2 = st.columns(2)
            with c_a1:
                s_act = st.text_input("Акт / Сюжетная арка:", value="Основной сюжет")
            with c_a2:
                s_status = st.selectbox("Статус готовности:", ["💡 Идея", "✍️ В работе", "✅ Готово"])
            
            s_content = st.text_area("Текст сценария (используйте #, ##, ### для скрытых блоков):", value=s_init_content, height=400)

            if st.button("💾 Сохранить сценарий", type="primary", use_container_width=True):
                if s_title.strip():
                    new_s = pd.DataFrame([{"Title": s_title.strip(), "Act": s_act.strip(), "Status": s_status, "Content": s_content}])
                    st.session_state['scripts'] = pd.concat([df_scr, new_s], ignore_index=True)
                    save_sheet("Scripts", st.session_state['scripts'])
                    st.session_state['selected_script'] = s_title.strip()
                    st.success("Сценарий сохранен!")
                    st.rerun()
                else:
                    st.warning("Введите название сценария!")

        elif st.session_state.get('selected_script') in df_scr['Title'].values:
            curr_s_title = st.session_state['selected_script']
            s_row = df_scr[df_scr['Title'] == curr_s_title].iloc[0]

            tab_s_read, tab_s_edit = st.tabs(["🎭 Интерактивное чтение (Со спойлерами и карточками)", "✏️ Редактировать / Загрузить из файла"])

            with tab_s_read:
                st.caption(f"Арка: **{s_row['Act']}** | Статус: **{s_row['Status']}**")
                st.markdown(f"## {s_row['Title']}")
                
                rendered_html, scr_found_chars = render_text_with_spoilers_and_characters(
                    str(s_row['Content']), CHAR_MATCHERS, CHAR_POPUP_SPANS, enable_spoilers=True
                )
                render_character_popover_bar(scr_found_chars, CHAR_LOOKUP, key_prefix="scr")
                st.write("---")
                st.caption("👆 Нажимайте на цветные блоки для раскрытия сцен, а на подсвеченные имена героев (`👤 Имя`) — для просмотра их карточки:")
                st.markdown(f"<div class='script-body'>{rendered_html}</div>", unsafe_allow_html=True)

            with tab_s_edit:
                rep_scr = st.file_uploader("📂 Обновить текст сценария из файла (.docx, .txt, .md):", type=["docx", "txt", "md"], key="scr_replace")
                s_curr_text = s_row['Content']
                if rep_scr is not None:
                    _, s_curr_text = read_uploaded_document(rep_scr)
                    st.info("Текст из файла подгружен. Нажмите «Сохранить изменения» внизу.")

                ed_s_title = st.text_input("Название:", value=s_row['Title'])
                ec1, ec2 = st.columns(2)
                with ec1:
                    ed_s_act = st.text_input("Арка / Акт:", value=s_row['Act'])
                with ec2:
                    statuses = ["💡 Идея", "✍️ В работе", "✅ Готово"]
                    curr_st_idx = statuses.index(s_row['Status']) if s_row['Status'] in statuses else 0
                    ed_s_status = st.selectbox("Статус:", statuses, index=curr_st_idx)

                ed_s_content = st.text_area("Текст сценария (# - Заголовок, ## - Подзаголовок, ### - Секрет):", value=s_curr_text, height=420)

                sb1, sb2 = st.columns([3, 1])
                with sb1:
                    if st.button("💾 Сохранить изменения", type="primary", use_container_width=True):
                        idx = df_scr[df_scr['Title'] == curr_s_title].index[0]
                        st.session_state['scripts'].at[idx, 'Title'] = ed_s_title.strip()
                        st.session_state['scripts'].at[idx, 'Act'] = ed_s_act.strip()
                        st.session_state['scripts'].at[idx, 'Status'] = ed_s_status
                        st.session_state['scripts'].at[idx, 'Content'] = ed_s_content
                        save_sheet("Scripts", st.session_state['scripts'])
                        st.session_state['selected_script'] = ed_s_title.strip()
                        st.success("Сценарий обновлен!")
                        st.rerun()
                with sb2:
                    if st.button("🗑 Удалить сценарий", use_container_width=True):
                        st.session_state['scripts'] = df_scr[df_scr['Title'] != curr_s_title].reset_index(drop=True)
                        save_sheet("Scripts", st.session_state['scripts'])
                        st.session_state['selected_script'] = None
                        st.rerun()


# ==============================================================================
# МОДУЛЬ 8: ЭКСПОРТ «БИБЛИИ ВСЕЛЕННОЙ» В WORD (.DOCX) ИЛИ MARKDOWN (.MD)
# ==============================================================================
elif app_mode == "📥 Экспорт «Библии Вселенной»":
    st.title("📥 Экспорт «Библии Вселенной» БИГ СИТИ")
    st.markdown("Соберите все разделы вашей базы данных в единый структурированный документ **Word (.docx)** или **Markdown (.md)** в один клик.")

    all_export_sections = [
        "Персонажи и Связи",
        "Игровые Карточки",
        "Фракции и Организации",
        "Карта Местности (Локации)",
        "Хронология Вселенной",
        "Архивы и Лор",
        "Сценарии и Сюжет"
    ]
    chosen_sections = st.multiselect(
        "Выберите разделы для включения в книгу:",
        all_export_sections,
        default=all_export_sections
    )

    compiled_md = generate_universe_markdown(
        chosen_sections,
        st.session_state['nodes'],
        st.session_state['edges'],
        st.session_state['cards'],
        st.session_state['factions'],
        st.session_state['locations'],
        st.session_state['timeline'],
        st.session_state['archives'],
        st.session_state['scripts']
    )
    compiled_docx = generate_universe_docx(compiled_md)

    ex_col1, ex_col2 = st.columns(2)
    with ex_col1:
        st.download_button(
            label="📘 Скачать «Библию Вселенной» в Word (.docx)",
            data=compiled_docx,
            file_name="BIG_CITY_Universe_Bible.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
            use_container_width=True
        )
    with ex_col2:
        st.download_button(
            label="📝 Скачать «Библию Вселенной» в Markdown (.md)",
            data=compiled_md.encode('utf-8'),
            file_name="BIG_CITY_Universe_Bible.md",
            mime="text/markdown",
            use_container_width=True
        )

    with st.expander("👁 Предпросмотр собираемого документа", expanded=False):
        st.markdown(compiled_md)
